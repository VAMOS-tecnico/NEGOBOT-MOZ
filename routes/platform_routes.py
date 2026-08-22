import csv
import base64
import hashlib
import hmac
import io
import logging
import os
import re
import secrets
import time
from datetime import datetime, timezone
from functools import wraps
from typing import Any, Callable
from urllib.parse import quote, urlparse

import requests
from flask import Blueprint, Response, jsonify, redirect, request, session, stream_with_context
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

import extensions
from services.trial_service import ACTIVE_STATUS, PENDING_STATUS, active_fields, is_expired as trial_is_expired, is_paid_plan, pending_fields
from services.central_account_service import central_account_id_for_tenant, claim_trial_for_account, pending_registry_fields, registry_is_expired, registry_status, trial_fields_from_registry
from services.channel_registry import CHANNEL_STATUSES, client_channel_rows, ensure_channel
from services.plan_service import entitlements_for_tenant, plan_channel_limit, public_plan_rows
from services.secret_store import SecretStoreError, decrypt_secret, encrypt_secret
from services.telegram_service import TelegramApiError, delete_webhook, get_me, get_webhook_info, set_webhook
from services.group_automation_service import archive_groups_for_instance, authorized_group_jids, group_document_id, sync_groups_for_tenant
from services.channel_publication_service import channel_capability, create_publication_data, enqueue_publication
from services.channel_oauth_service import complete_oauth, disconnect_oauth, provider_config, start_oauth
from services.password_reset_service import consume_password_reset, request_password_reset
from services.ai_queue_service import AIQueueError, request_ai_text
from services.evolution_service import get_connection_state, get_profile_picture_url, listar_chats_whatsapp, send_media, send_whatsapp
from services.knowledge_base_service import KnowledgeBaseError, build_tenant_context, delete_original, extract_text, list_tenant_files, read_blob, serialise_file, store_blob, store_original, validate_upload

logger = logging.getLogger(__name__)
platform_bp = Blueprint("platform", __name__, url_prefix="/api/platform")
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_LOGIN_ATTEMPTS: dict[str, list[float]] = {}
_LOGIN_WINDOW_SECONDS = 900
_LOGIN_MAX_ATTEMPTS = 8
_VIDEO_MAX_TOTAL_SCRIPT_CHARACTERS = 5_000
_VIDEO_MAX_TOTAL_DURATION_SECONDS = 300


def _request_key(identifier: str) -> str:
    forwarded = request.headers.get("X-Forwarded-For", "").split(",")[0].strip()
    address = forwarded or request.remote_addr or "unknown"
    return f"{address}:{identifier[:160]}"


def _login_allowed(identifier: str) -> bool:
    now = time.time()
    key = _request_key(identifier)
    recent = [stamp for stamp in _LOGIN_ATTEMPTS.get(key, []) if now - stamp < _LOGIN_WINDOW_SECONDS]
    if len(recent) >= _LOGIN_MAX_ATTEMPTS:
        _LOGIN_ATTEMPTS[key] = recent
        return False
    recent.append(now)
    _LOGIN_ATTEMPTS[key] = recent
    return True


@platform_bp.before_request
def _same_origin_mutation_guard():
    if request.method not in {"POST", "PUT", "PATCH", "DELETE"}:
        return None
    if request.path in {"/api/platform/auth/login", "/api/platform/auth/logout"}:
        return None
    origin = request.headers.get("Origin")
    if origin:
        origin_host = urlparse(origin).netloc
        if origin_host and origin_host != request.host:
            return jsonify({"error": "Origem da operação não autorizada."}), 403
    return None


def _now():
    return datetime.now(timezone.utc)


def _plan_expired(data: dict[str, Any], now: datetime | None = None) -> bool:
    """Expira trial apenas após ligação WhatsApp confirmada; preserva planos pagos."""
    return trial_is_expired(data, now)


def _db():
    if extensions.db is None:
        raise RuntimeError("Base de dados da plataforma indisponível")
    return extensions.db


def _doc_id(value: str) -> str:
    return hashlib.sha256(value.strip().lower().encode("utf-8")).hexdigest()


def _identity() -> dict[str, Any] | None:
    value = session.get("platform_identity")
    return value if isinstance(value, dict) and value.get("role") else None


def _require_roles(*roles: str) -> Callable:
    def decorator(handler: Callable) -> Callable:
        @wraps(handler)
        def wrapped(*args, **kwargs):
            identity = _identity()
            if not identity:
                return jsonify({"error": "autenticação necessária"}), 401
            if identity.get("role") not in roles:
                return jsonify({"error": "permissão insuficiente"}), 403
            return handler(*args, **kwargs)

        return wrapped

    return decorator


def _require_tenant_roles(*tenant_roles: str) -> Callable:
    def decorator(handler: Callable) -> Callable:
        @wraps(handler)
        def wrapped(*args, **kwargs):
            identity = _identity()
            if not identity:
                return jsonify({"error": "autenticação necessária"}), 401
            if identity.get("role") not in {"client", "operator"}:
                return jsonify({"error": "A operação exige uma conta de tenant."}), 403
            if identity.get("tenant_role") not in tenant_roles:
                return jsonify({"error": "A tua função não permite esta operação."}), 403
            return handler(*args, **kwargs)

        return wrapped

    return decorator


def _tenant_for_identity(identity: dict[str, Any]) -> str | None:
    tenant_id = identity.get("tenant_id")
    return str(tenant_id) if tenant_id else None


def _create_platform_account(
    db: Any,
    *,
    name: str,
    email: str,
    password: str,
    billing_region: str,
    selected_plan: str,
) -> dict[str, Any] | None:
    """Create the account, tenant and trial registry as one uniqueness boundary.

    The canonical platform user document is keyed by the normalized email hash.
    In production, a Firestore transaction makes two simultaneous registrations
    for the same email conflict instead of creating two tenants. Test doubles
    without transaction support use the same deterministic document check.
    """
    user_ref = db.collection("platform_users").document(_doc_id(email))
    central_account_id = f"ca_{user_ref.id[:24]}"
    registry_ref = db.collection("central_trial_registry").document(central_account_id)
    tenant_id = f"tnt_{secrets.token_urlsafe(8)}"
    tenant_ref = db.collection("tenants").document(tenant_id)
    now = _now()
    password_hash = generate_password_hash(password)
    tenant_fields = {
        "name": name,
        "email": email,
        "account_email": email,
        "central_account_id": central_account_id,
        "central_identity_email_hash": _doc_id(email),
        "empresa_nome": name,
        "status": "active",
        "plan": "demonstracao",
        "plano": "demonstracao",
        "nome_plano": "Demonstração",
        "status_plano": "demonstracao",
        "trial_status": "trial_pending_connection",
        "trial_connection_confirmed": False,
        "billing_region": billing_region,
        "selected_plan": selected_plan or None,
        "central_trial_status": PENDING_STATUS,
        "onboarding_status": "incomplete",
        "profile_completed": False,
        "onboarding_step": "profile",
        "created_at": now,
        "limits": {"contacts": 500, "contact_limit": 500, "conversation_limit": 500, "campaigns_per_month": 2, "team_seats": 1, "messages_per_campaign": 100, "included_channels": ["whatsapp"], "additional_channel_slots": 0},
    }
    user_fields = {
        "name": name,
        "email": email,
        "central_account_id": central_account_id,
        "role": "client",
        "tenant_id": tenant_id,
        "tenant_role": "owner",
        "status": "active",
        "password_hash": password_hash,
        "created_at": now,
        "last_login_at": now,
    }
    registry_fields = pending_registry_fields(central_account_id, tenant_id, email=email, now=now)

    transaction_factory = getattr(db, "transaction", None)
    if not callable(transaction_factory):
        if user_ref.get().exists or registry_ref.get().exists:
            return None
        tenant_ref.set(tenant_fields)
        user_ref.set(user_fields)
        registry_ref.set(registry_fields, merge=True)
        return {"tenant_id": tenant_id, "central_account_id": central_account_id, "now": now}

    last_error: Exception | None = None
    for _attempt in range(3):
        try:
            transaction = transaction_factory()
            existing_user = transaction.get(user_ref)
            existing_registry = transaction.get(registry_ref)
            if getattr(existing_user, "exists", False) or getattr(existing_registry, "exists", False):
                return None
            transaction.set(tenant_ref, tenant_fields)
            transaction.set(user_ref, user_fields)
            transaction.set(registry_ref, registry_fields, merge=True)
            transaction.commit()
            return {"tenant_id": tenant_id, "central_account_id": central_account_id, "now": now}
        except Exception as exc:
            last_error = exc
    raise RuntimeError("Não foi possível criar a conta de forma segura.") from last_error


def _tenant_data(tenant_id: str | None) -> dict[str, Any]:
    if not tenant_id:
        return {}
    document = _db().collection("tenants").document(tenant_id).get()
    return document.to_dict() if document.exists else {}


def _current_entitlements(tenant_id: str | None) -> dict[str, Any]:
    return entitlements_for_tenant(_tenant_data(tenant_id))


def _month_key(value: Any) -> str:
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m")
    return str(value or "")[:7]


def _count_tenant_campaigns_this_month(tenant_id: str) -> int:
    month = _now().strftime("%Y-%m")
    return sum(
        1
        for document in _db().collection("campaigns").where("tenant_id", "==", tenant_id).limit(5000).stream()
        if _month_key((document.to_dict() or {}).get("created_at")) == month
    )


def _count_tenant_contacts(tenant_id: str, limit: int = 20000) -> int:
    return sum(1 for _ in _db().collection("contacts").where("tenant_id", "==", tenant_id).limit(limit).stream())


def _count_tenant_team(tenant_id: str, limit: int = 100) -> int:
    return sum(1 for _ in _db().collection("platform_users").where("tenant_id", "==", tenant_id).limit(limit).stream())


@platform_bp.post("/auth/login")
def login():
    payload = request.get_json(silent=True) or {}
    identifier = str(payload.get("email") or payload.get("identifier") or "").strip().lower()
    password = str(payload.get("password") or "")
    if not identifier or not password:
        return jsonify({"error": "Introduza o email e a palavra-passe."}), 400
    if not _login_allowed(identifier):
        return jsonify({"error": "Demasiadas tentativas. Aguarda alguns minutos antes de tentar novamente."}), 429

    admin_token = os.getenv("ADMIN_TOKEN", "")
    if identifier in {"admin", "owner", "administrador"} and admin_token and hmac.compare_digest(password, admin_token):
        identity = {"id": "platform-owner", "name": "Administrador", "role": "owner", "tenant_id": None}
        session.clear()
        session["platform_identity"] = identity
        session.permanent = True
        return jsonify({"authenticated": True, "user": identity})

    if not _EMAIL_RE.fullmatch(identifier) or len(password) < 8:
        return jsonify({"error": "Credenciais inválidas."}), 401
    try:
        document = _db().collection("platform_users").document(_doc_id(identifier)).get()
    except Exception:
        return jsonify({"error": "Não foi possível consultar a conta neste momento."}), 503
    if not document.exists:
        return jsonify({"error": "Credenciais inválidas."}), 401
    user = document.to_dict() or {}
    if user.get("status", "active") != "active" or not check_password_hash(user.get("password_hash", ""), password):
        return jsonify({"error": "Credenciais inválidas."}), 401
    identity = {
        "id": document.id,
        "name": user.get("name") or identifier,
        "email": identifier,
        "role": user.get("role", "client"),
        "tenant_id": user.get("tenant_id"),
        "tenant_role": user.get("tenant_role"),
    }
    session.clear()
    session["platform_identity"] = identity
    session.permanent = True
    document.reference.set({"last_login_at": _now()}, merge=True)
    return jsonify({"authenticated": True, "user": identity})


@platform_bp.post("/auth/forgot-password")
def forgot_password():
    payload = request.get_json(silent=True) or {}
    email = str(payload.get("email") or "").strip().lower()
    frontend = str(os.getenv("PUBLIC_APP_BASE_URL") or "https://app-negobotmoz.duckdns.org/plataforma").rstrip("/")
    try:
        request_password_reset(_db(), email, frontend)
    except Exception:
        # Keep the same public response for unknown accounts and delivery errors.
        pass
    return jsonify({
        "accepted": True,
        "message_en": "If an account exists with this email, you will receive a password reset link.",
        "message_pt": "Se existir uma conta com esse email, receberás uma ligação para trocar a palavra-passe.",
    })


@platform_bp.post("/auth/reset-password")
def reset_password():
    payload = request.get_json(silent=True) or {}
    token = str(payload.get("token") or "").strip()
    password = str(payload.get("password") or "")
    if len(password) < 8:
        return jsonify({"error": "A nova palavra-passe deve ter pelo menos 8 caracteres."}), 400
    if not consume_password_reset(_db(), token, password):
        return jsonify({"error": "A ligação é inválida, já foi usada ou expirou."}), 400
    return jsonify({
        "reset": True,
        "message_en": "Password changed. You can now sign in to the platform.",
        "message_pt": "Palavra-passe alterada. Já podes entrar na plataforma.",
    })


@platform_bp.post("/auth/register")
def register():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or "")
    billing_region = str(payload.get("billing_region") or "mozambique").strip().lower()
    selected_plan = str(payload.get("plan_id") or "").strip().lower()
    if not _EMAIL_RE.fullmatch(email) or len(password) < 8:
        return jsonify({"error": "Email válido e palavra-passe com pelo menos 8 caracteres são obrigatórios."}), 400
    if not name:
        name = email.split("@", 1)[0].replace(".", " ").replace("_", " ").strip() or "Novo cliente"
    if len(name) < 2:
        name = "Novo cliente"
    if billing_region not in {"mozambique", "international"}:
        return jsonify({"error": "Escolhe Moçambique ou pagamento internacional."}), 400
    if selected_plan and selected_plan not in {"basico", "medio", "premium"}:
        return jsonify({"error": "Plano selecionado inválido."}), 400
    if not _login_allowed(email):
        return jsonify({"error": "Demasiadas tentativas. Aguarda alguns minutos antes de tentar novamente."}), 429
    try:
        created = _create_platform_account(
            _db(),
            name=name,
            email=email,
            password=password,
            billing_region=billing_region,
            selected_plan=selected_plan,
        )
    except RuntimeError:
        return jsonify({"error": "Não foi possível criar a conta de forma segura neste momento."}), 503
    if created is None:
        return jsonify({"error": "Já existe uma conta com este email. Usa a opção Entrar na plataforma."}), 409
    tenant_id = str(created["tenant_id"])
    central_account_id = str(created["central_account_id"])
    user_ref = _db().collection("platform_users").document(_doc_id(email))
    identity = {"id": user_ref.id, "name": name, "email": email, "central_account_id": central_account_id, "role": "client", "tenant_id": tenant_id, "tenant_role": "owner"}
    session.clear()
    session["platform_identity"] = identity
    session.permanent = True
    return jsonify({"authenticated": True, "user": identity, "tenant": {"id": tenant_id, "name": name, "plan": "demonstracao"}}), 201


@platform_bp.post("/auth/logout")
def logout():
    session.pop("platform_identity", None)
    return jsonify({"authenticated": False})


@platform_bp.get("/auth/me")
def me():
    identity = _identity()
    return jsonify({"authenticated": bool(identity), "user": identity})


@platform_bp.get("/admin/overview")
@_require_roles("owner", "admin")
def admin_overview():
    db = _db()
    tenants = list(db.collection("tenants").stream())
    users = list(db.collection("platform_users").stream())
    return jsonify({
        "role": _identity().get("role"),
        "tenants": len(tenants),
        "users": len(users),
        "active_tenants": sum(1 for item in tenants if (item.to_dict() or {}).get("status", "active") == "active"),
        "features": ["clientes", "instâncias", "campanhas", "integrações", "pagamentos", "auditoria"],
    })


@platform_bp.get("/admin/tenants")
@_require_roles("owner", "admin")
def list_tenants():
    rows = []
    for document in _db().collection("tenants").stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    rows.sort(key=lambda item: item.get("created_at", ""), reverse=True)
    return jsonify({"tenants": rows})


@platform_bp.post("/admin/tenants")
@_require_roles("owner", "admin")
def create_tenant():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or "")
    if len(name) < 2 or not _EMAIL_RE.fullmatch(email) or len(password) < 8:
        return jsonify({"error": "Nome, email válido e palavra-passe com pelo menos 8 caracteres são obrigatórios."}), 400
    db = _db()
    user_ref = db.collection("platform_users").document(_doc_id(email))
    if user_ref.get().exists:
        return jsonify({"error": "Já existe uma conta com este email."}), 409
    tenant_id = f"tnt_{secrets.token_urlsafe(8)}"
    central_account_id = f"ca_{user_ref.id[:24]}"
    now = _now()
    tenant_ref = db.collection("tenants").document(tenant_id)
    tenant_ref.set({
        "name": name,
        "email": email,
        "account_email": email,
        "central_account_id": central_account_id,
        "central_identity_email_hash": _doc_id(email),
        "empresa_nome": name,
        "status": "active",
        "plan": "demonstracao",
        "plano": "demonstracao",
        "nome_plano": "Demonstração",
        "status_plano": "demonstracao",
        "trial_status": "trial_pending_connection",
        "trial_connection_confirmed": False,
        "created_at": now,
        "limits": {"contacts": 500, "contact_limit": 500, "conversation_limit": 500, "campaigns_per_month": 2, "team_seats": 1, "messages_per_campaign": 100, "included_channels": ["whatsapp"], "additional_channel_slots": 0},
    })
    user_ref.set({
        "name": name,
        "email": email,
        "central_account_id": central_account_id,
        "role": "client",
        "tenant_id": tenant_id,
        "tenant_role": "owner",
        "status": "active",
        "password_hash": generate_password_hash(password),
        "created_at": now,
        "last_login_at": None,
    })
    db.collection("central_trial_registry").document(central_account_id).set(pending_registry_fields(central_account_id, tenant_id, email=email, now=now), merge=True)
    return jsonify({"created": True, "tenant": {"id": tenant_id, "name": name, "plan": "demonstracao"}}), 201


@platform_bp.get("/client/team")
@_require_roles("client", "operator")
def list_tenant_team():
    tenant_id = _tenant_for_identity(_identity())
    rows = []
    for document in _db().collection("platform_users").where("tenant_id", "==", tenant_id).limit(100).stream():
        item = document.to_dict() or {}
        rows.append({
            "id": document.id,
            "name": item.get("name", ""),
            "email": item.get("email", ""),
            "role": item.get("role", "client"),
            "tenant_role": item.get("tenant_role", "viewer"),
            "status": item.get("status", "active"),
            "created_at": item.get("created_at"),
            "last_login_at": item.get("last_login_at"),
        })
    rows.sort(key=lambda item: item.get("name", "").lower())
    return jsonify({"users": rows, "current_role": (_identity() or {}).get("tenant_role")})


@platform_bp.post("/client/team")
@_require_tenant_roles("owner")
def create_tenant_operator():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    email = str(payload.get("email") or "").strip().lower()
    password = str(payload.get("password") or "")
    if len(name) < 2 or not _EMAIL_RE.fullmatch(email) or len(password) < 8:
        return jsonify({"error": "Nome, email válido e palavra-passe com pelo menos 8 caracteres são obrigatórios."}), 400
    db = _db()
    user_ref = db.collection("platform_users").document(_doc_id(email))
    if user_ref.get().exists:
        return jsonify({"error": "Já existe uma conta com este email."}), 409
    tenant_id = _tenant_for_identity(_identity())
    entitlements = _current_entitlements(tenant_id)
    team_limit = int(entitlements.get("team_seats") or 1)
    if _count_tenant_team(tenant_id or "") >= team_limit:
        return jsonify({"error": f"O teu plano permite até {team_limit} utilizador(es). Faz upgrade para adicionar mais membros."}), 403
    now = _now()
    user_ref.set({
        "name": name,
        "email": email,
        "role": "operator",
        "tenant_role": "operator",
        "tenant_id": tenant_id,
        "status": "active",
        "password_hash": generate_password_hash(password),
        "created_at": now,
        "last_login_at": None,
    })
    _audit("tenant_operator_created", _identity(), tenant_id, {"user_id": user_ref.id})
    return jsonify({"created": True, "user": {"id": user_ref.id, "name": name, "email": email, "role": "operator", "tenant_role": "operator", "status": "active"}}), 201


@platform_bp.patch("/client/team/<user_id>")
@_require_tenant_roles("owner")
def update_tenant_operator(user_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("platform_users").document(user_id)
    document = reference.get()
    data = document.to_dict() if document.exists else None
    if not data or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Utilizador não encontrado neste tenant."}), 404
    if data.get("tenant_role") == "owner":
        return jsonify({"error": "O proprietário principal não pode ser alterado por esta operação."}), 400
    payload = request.get_json(silent=True) or {}
    allowed = {}
    if "status" in payload:
        status = str(payload.get("status") or "").lower()
        if status not in {"active", "suspended"}:
            return jsonify({"error": "Estado inválido."}), 400
        allowed["status"] = status
    if "tenant_role" in payload:
        role = str(payload.get("tenant_role") or "").lower()
        if role not in {"operator", "viewer"}:
            return jsonify({"error": "Função inválida."}), 400
        allowed["tenant_role"] = role
        allowed["role"] = "operator" if role == "operator" else "client"
    if not allowed:
        return jsonify({"error": "Nenhuma alteração válida foi enviada."}), 400
    allowed["updated_at"] = _now()
    reference.set(allowed, merge=True)
    _audit("tenant_operator_updated", _identity(), tenant_id, {"user_id": user_id, "fields": sorted(allowed)})
    return jsonify({"updated": True, "user_id": user_id, "fields": sorted(allowed)})


@platform_bp.get("/client/overview")
@_require_roles("client", "operator")
def client_overview():
    identity = _identity()
    tenant_id = _tenant_for_identity(identity)
    if not tenant_id:
        return jsonify({"error": "tenant não configurado"}), 403
    db = _db()
    tenant = db.collection("tenants").document(tenant_id).get()
    if not tenant.exists:
        return jsonify({"error": "tenant não encontrado"}), 404
    contacts = list(db.collection("contacts").where("tenant_id", "==", tenant_id).limit(500).stream())
    campaigns = list(db.collection("campaigns").where("tenant_id", "==", tenant_id).limit(100).stream())
    return jsonify({
        "role": identity.get("role"),
        "tenant": {"id": tenant_id, **(tenant.to_dict() or {})},
        "contacts": len(contacts),
        "campaigns": len(campaigns),
        "conversations": 0,
        "features": ["contactos", "segmentos", "campanhas", "disparos", "conversas", "assistente", "integração WhatsApp"],
    })


_SOCIAL_PROFILE_KEYS = ("facebook", "instagram", "twitter_x", "tiktok", "telegram", "linkedin")


@platform_bp.get("/client/profile")
@_require_roles("client", "operator")
def get_client_profile():
    identity = _identity() or {}
    tenant_id = _tenant_for_identity(identity)
    tenant = _db().collection("tenants").document(tenant_id).get().to_dict() or {}
    socials = tenant.get("redes_sociais") if isinstance(tenant.get("redes_sociais"), dict) else {}
    return jsonify({
        "email": tenant.get("account_email") or tenant.get("email") or identity.get("email"),
        "empresa_nome": tenant.get("empresa_nome") or tenant.get("name", ""),
        "nicho": tenant.get("nicho", ""),
        "email_corporativo": tenant.get("email_corporativo", ""),
        "redes_sociais": {key: str(socials.get(key) or "") for key in _SOCIAL_PROFILE_KEYS},
        "instance_name": tenant.get("instance_name"),
        "status_conexao": tenant.get("evolution_state", "desconectado"),
        "billing_region": tenant.get("billing_region", "mozambique"),
        "selected_plan": tenant.get("selected_plan"),
        "preferred_trial_channel": tenant.get("preferred_trial_channel", "whatsapp"),
        "onboarding_status": tenant.get("onboarding_status", "incomplete"),
        "profile_completed": bool(tenant.get("profile_completed", False)),
    })


@platform_bp.patch("/client/profile")
@_require_tenant_roles("owner", "operator")
def update_client_profile():
    payload = request.get_json(silent=True) or {}
    tenant_id = _tenant_for_identity(_identity())
    changes = {}
    if "empresa_nome" in payload:
        value = str(payload.get("empresa_nome") or "").strip()[:160]
        if len(value) < 2:
            return jsonify({"error": "Indica o nome da empresa."}), 400
        changes["empresa_nome"] = value
        changes["name"] = value
    if "nicho" in payload:
        changes["nicho"] = str(payload.get("nicho") or "").strip()[:160]
    if "email_corporativo" in payload:
        email = str(payload.get("email_corporativo") or "").strip().lower()
        if email and not _EMAIL_RE.fullmatch(email):
            return jsonify({"error": "Indica um email corporativo válido."}), 400
        changes["email_corporativo"] = email
    if "redes_sociais" in payload:
        if not isinstance(payload.get("redes_sociais"), dict):
            return jsonify({"error": "As redes sociais devem ser enviadas como objeto."}), 400
        changes["redes_sociais"] = {
            key: str(payload["redes_sociais"].get(key) or "").strip()[:400]
            for key in _SOCIAL_PROFILE_KEYS
        }
    if "billing_region" in payload:
        region = str(payload.get("billing_region") or "").strip().lower()
        if region not in {"mozambique", "international"}:
            return jsonify({"error": "Escolhe Moçambique ou pagamento internacional."}), 400
        changes["billing_region"] = region
    if "selected_plan" in payload or "plan_id" in payload:
        selected_plan = str(payload.get("selected_plan") or payload.get("plan_id") or "").strip().lower()
        if selected_plan and selected_plan not in {"basico", "medio", "premium"}:
            return jsonify({"error": "Plano seleccionado inválido."}), 400
        changes["selected_plan"] = selected_plan or None
    if "preferred_trial_channel" in payload:
        preferred_channel = str(payload.get("preferred_trial_channel") or "").strip().lower()
        if preferred_channel not in {"whatsapp", "telegram", "instagram", "facebook"}:
            return jsonify({"error": "Canal de teste inválido."}), 400
        changes["preferred_trial_channel"] = preferred_channel
    if not changes:
        return jsonify({"error": "Nenhuma alteração de perfil foi enviada."}), 400
    changes["updated_at"] = _now()
    tenant_ref = _db().collection("tenants").document(tenant_id)
    current_tenant = tenant_ref.get().to_dict() or {}
    effective_company = changes.get("empresa_nome", current_tenant.get("empresa_nome") or current_tenant.get("name"))
    effective_region = changes.get("billing_region", current_tenant.get("billing_region", "mozambique"))
    if str(effective_company or "").strip() and effective_region in {"mozambique", "international"}:
        changes["profile_completed"] = True
        changes["onboarding_status"] = "completed"
        changes["onboarding_step"] = "whatsapp"
    tenant_ref.set(changes, merge=True)
    current_tenant = {**current_tenant, **changes}
    instance_name = str(current_tenant.get("instance_name") or "").strip()
    if instance_name:
        _db().collection("clientes_bot").document(instance_name).set({
            key: changes[key] for key in ("empresa_nome", "nicho", "email_corporativo", "redes_sociais") if key in changes
        }, merge=True)
    _audit("client_profile_updated", _identity(), tenant_id, {"fields": sorted(changes)})
    return jsonify({"updated": True, "fields": sorted(changes)})


@platform_bp.get("/client/contacts")
@_require_roles("client", "operator")
def list_contacts():
    tenant_id = _tenant_for_identity(_identity())
    search = str(request.args.get("search") or "").strip().lower()[:120]
    tag = str(request.args.get("tag") or "").strip().lower()[:80]
    opt_in_filter = str(request.args.get("opt_in") or "").strip().lower()
    query = _db().collection("contacts").where("tenant_id", "==", tenant_id)
    if tag:
        query = query.where("tags", "array_contains", tag)
    if opt_in_filter in {"true", "false"}:
        query = query.where("opt_in", "==", opt_in_filter == "true")
    rows = []
    for document in query.limit(1000).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        item.setdefault("tags", [])
        if search and search not in f"{item.get('name', '')} {item.get('phone', '')}".lower():
            continue
        if item.get("status", "active") == "archived":
            continue
        rows.append(item)
    rows.sort(key=lambda item: str(item.get("name") or "").lower())
    return jsonify({"contacts": rows, "count": len(rows), "filters": {"search": search, "tag": tag, "opt_in": opt_in_filter or None}})


@platform_bp.post("/client/contacts")
@_require_roles("client", "operator")
def create_contact():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    phone = re.sub(r"\D", "", str(payload.get("phone") or ""))
    if len(name) < 2 or len(phone) < 8:
        return jsonify({"error": "Nome e número de telefone válidos são obrigatórios."}), 400
    tenant_id = _tenant_for_identity(_identity())
    entitlements = _current_entitlements(tenant_id)
    contact_limit = int(entitlements.get("contact_limit") or 0)
    if contact_limit and _count_tenant_contacts(tenant_id or "", contact_limit + 1) >= contact_limit:
        return jsonify({"error": f"O teu plano permite até {contact_limit} contactos. Faz upgrade para adicionar mais."}), 403
    tags = sorted({str(item).strip().lower() for item in (payload.get("tags") or []) if str(item).strip()})[:20]
    ref = _db().collection("contacts").document()
    ref.set({"tenant_id": tenant_id, "name": name, "phone": phone, "opt_in": bool(payload.get("opt_in", True)), "tags": tags, "status": "active", "created_at": _now(), "updated_at": _now()})
    _audit("contact_created", _identity(), tenant_id, {"contact_id": ref.id})
    return jsonify({"created": True, "contact": {"id": ref.id, "name": name, "phone": phone, "opt_in": bool(payload.get("opt_in", True)), "tags": tags}}), 201


@platform_bp.patch("/client/contacts/<contact_id>")
@_require_roles("client", "operator")
def update_contact(contact_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("contacts").document(contact_id)
    document = reference.get()
    data = document.to_dict() if document.exists else None
    if not data or data.get("tenant_id") != tenant_id or data.get("status", "active") == "archived":
        return jsonify({"error": "Contacto não encontrado neste tenant."}), 404
    payload = request.get_json(silent=True) or {}
    changes = {}
    if "name" in payload:
        name = str(payload.get("name") or "").strip()
        if len(name) < 2:
            return jsonify({"error": "Nome de contacto inválido."}), 400
        changes["name"] = name
    if "phone" in payload:
        phone = re.sub(r"\D", "", str(payload.get("phone") or ""))
        if len(phone) < 8:
            return jsonify({"error": "Número de contacto inválido."}), 400
        changes["phone"] = phone
    if "opt_in" in payload:
        changes["opt_in"] = bool(payload.get("opt_in"))
    if "tags" in payload:
        if not isinstance(payload.get("tags"), list):
            return jsonify({"error": "As etiquetas devem ser uma lista."}), 400
        changes["tags"] = sorted({str(item).strip().lower() for item in payload["tags"] if str(item).strip()})[:20]
    if not changes:
        return jsonify({"error": "Nenhuma alteração válida foi enviada."}), 400
    changes["updated_at"] = _now()
    reference.set(changes, merge=True)
    _audit("contact_updated", _identity(), tenant_id, {"contact_id": contact_id, "fields": sorted(changes)})
    return jsonify({"updated": True, "contact_id": contact_id, "fields": sorted(changes)})


@platform_bp.delete("/client/contacts/<contact_id>")
@_require_roles("client", "operator")
def archive_contact(contact_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("contacts").document(contact_id)
    document = reference.get()
    data = document.to_dict() if document.exists else None
    if not data or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Contacto não encontrado neste tenant."}), 404
    reference.set({"status": "archived", "opt_in": False, "updated_at": _now()}, merge=True)
    _audit("contact_archived", _identity(), tenant_id, {"contact_id": contact_id})
    return jsonify({"archived": True, "contact_id": contact_id})


@platform_bp.get("/client/templates")
@_require_roles("client", "operator")
def list_campaign_templates():
    tenant_id = _tenant_for_identity(_identity())
    rows = []
    for document in _db().collection("campaign_templates").where("tenant_id", "==", tenant_id).limit(100).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    rows.sort(key=lambda item: str(item.get("name") or "").lower())
    return jsonify({"templates": rows})


@platform_bp.post("/client/templates")
@_require_roles("client", "operator")
def create_campaign_template():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    body = str(payload.get("body") or payload.get("message") or "").strip()
    if not 2 <= len(name) <= 100 or not 1 <= len(body) <= 4000:
        return jsonify({"error": "O nome deve ter 2–100 caracteres e a mensagem 1–4000 caracteres."}), 400
    variables = sorted({str(item).strip().lower() for item in (payload.get("variables") or []) if str(item).strip()})[:30]
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("campaign_templates").document()
    reference.set({"tenant_id": tenant_id, "name": name, "body": body, "variables": variables, "status": "active", "created_at": _now(), "updated_at": _now()})
    _audit("campaign_template_created", _identity(), tenant_id, {"template_id": reference.id})
    return jsonify({"created": True, "template": {"id": reference.id, "name": name, "body": body, "variables": variables, "status": "active"}}), 201


@platform_bp.patch("/client/templates/<template_id>")
@_require_roles("client", "operator")
def update_campaign_template(template_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("campaign_templates").document(template_id)
    document = reference.get()
    data = document.to_dict() if document.exists else None
    if not data or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Template não encontrado."}), 404
    payload = request.get_json(silent=True) or {}
    changes = {}
    if "name" in payload:
        name = str(payload.get("name") or "").strip()
        if not 2 <= len(name) <= 100:
            return jsonify({"error": "Nome de template inválido."}), 400
        changes["name"] = name
    if "body" in payload:
        body = str(payload.get("body") or "").strip()
        if not 1 <= len(body) <= 4000:
            return jsonify({"error": "Mensagem de template inválida."}), 400
        changes["body"] = body
    if "status" in payload:
        status = str(payload.get("status") or "").lower()
        if status not in {"active", "archived"}:
            return jsonify({"error": "Estado de template inválido."}), 400
        changes["status"] = status
    if not changes:
        return jsonify({"error": "Nenhuma alteração válida foi enviada."}), 400
    changes["updated_at"] = _now()
    reference.set(changes, merge=True)
    _audit("campaign_template_updated", _identity(), tenant_id, {"template_id": template_id, "fields": sorted(changes)})
    return jsonify({"updated": True, "template_id": template_id, "fields": sorted(changes)})


@platform_bp.get("/client/groups")
@_require_roles("client", "operator")
def client_groups():
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    connection_state = get_connection_state(instance_name) if instance_name else "not_configured"
    if connection_state != "open":
        archived = archive_groups_for_instance(instance_name) if instance_name else 0
        return jsonify({"tenant_id": tenant_id, "groups": [], "connection_state": connection_state, "archived": archived})
    documents = _db().collection("whatsapp_groups").where("tenant_id", "==", tenant_id).limit(500).stream()
    rows = []
    for document in documents:
        data = document.to_dict() or {}
        if data.get("status") == "archived" or data.get("visible") is False:
            continue
        data["id"] = document.id
        rows.append(data)
    rows.sort(key=lambda item: str(item.get("name") or item.get("group_jid") or "").lower())
    return jsonify({"tenant_id": tenant_id, "groups": rows, "connection_state": connection_state})


@platform_bp.post("/client/groups/sync")
@_require_tenant_roles("owner", "operator")
def sync_client_groups():
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    if not instance_name:
        return jsonify({"error": "Liga primeiro o WhatsApp deste tenant."}), 409
    try:
        result = sync_groups_for_tenant(tenant_id, instance_name)
        return jsonify(result)
    except PermissionError as exc:
        return jsonify({"error": str(exc)}), 403
    except requests.RequestException:
        return jsonify({"error": "A Evolution API não respondeu à sincronização de grupos."}), 502
    except Exception:
        return jsonify({"error": "Não foi possível sincronizar os grupos deste tenant."}), 503


@platform_bp.patch("/client/groups/<group_id>")
@_require_tenant_roles("owner", "operator")
def update_client_group(group_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("whatsapp_groups").document(group_id)
    document = reference.get()
    data = document.to_dict() if document.exists else {}
    if not document.exists or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Grupo não encontrado."}), 404
    payload = request.get_json(silent=True) or {}
    allowed = {"automation_enabled", "mention_required", "welcome_enabled", "welcome_message", "keywords"}
    changes = {key: payload[key] for key in allowed if key in payload}
    if not changes:
        return jsonify({"error": "Nenhuma configuração reconhecida."}), 400
    if changes.get("automation_enabled") and not (data.get("admin_verified") and data.get("bot_is_admin") and data.get("status") == "active"):
        return jsonify({"error": "Só podes activar automação num grupo onde a instância esteja verificada como administradora."}), 403
    if "welcome_message" in changes:
        changes["welcome_message"] = str(changes["welcome_message"] or "").strip()[:1000]
    if "mention_required" in changes:
        changes["mention_required"] = bool(changes["mention_required"])
    if "welcome_enabled" in changes:
        changes["welcome_enabled"] = bool(changes["welcome_enabled"])
    if "automation_enabled" in changes:
        changes["automation_enabled"] = bool(changes["automation_enabled"])
    if "keywords" in changes:
        raw_keywords = changes["keywords"]
        if not isinstance(raw_keywords, list) or len(raw_keywords) > 30:
            return jsonify({"error": "Define até 30 keywords."}), 400
        normalized_keywords = []
        for item in raw_keywords:
            if not isinstance(item, dict):
                continue
            trigger = str(item.get("trigger") or item.get("keyword") or "").strip()[:80]
            response = str(item.get("response") or item.get("text") or "").strip()[:1500]
            if trigger and response:
                normalized_keywords.append({"trigger": trigger, "response": response})
        changes["keywords"] = normalized_keywords
    reference.set({**changes, "updated_at": _now()}, merge=True)
    return jsonify({"updated": True, "group_id": group_id, "changes": changes})


@platform_bp.get("/admin/groups")
@_require_roles("owner", "admin")
def admin_groups():
    documents = _db().collection("whatsapp_groups").limit(1000).stream()
    rows = []
    for document in documents:
        data = document.to_dict() or {}
        data["id"] = document.id
        rows.append(data)
    rows.sort(key=lambda item: str(item.get("last_synced_at") or 0), reverse=True)
    return jsonify({"groups": rows})


@platform_bp.get("/client/whatsapp-channels/capability")
@_require_roles("client", "operator")
def client_whatsapp_channel_capability():
    return jsonify(channel_capability())


@platform_bp.get("/client/channel-publications")
@_require_roles("client", "operator")
def list_channel_publications():
    tenant_id = _tenant_for_identity(_identity())
    rows = []
    for document in _db().collection("channel_publications").where("tenant_id", "==", tenant_id).limit(200).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    rows.sort(key=lambda item: str(item.get("created_at") or ""), reverse=True)
    return jsonify({"publications": rows, "capability": channel_capability()})


@platform_bp.post("/client/channel-publications")
@_require_tenant_roles("owner", "operator")
def create_channel_publication():
    tenant_id = _tenant_for_identity(_identity())
    entitlements = _current_entitlements(tenant_id)
    if not entitlements.get("mass_broadcast"):
        return jsonify({"error": "As publicações agendadas estão disponíveis durante o trial Premium ou num plano Premium activo."}), 403
    payload = request.get_json(silent=True) or {}
    try:
        data = create_publication_data(payload, tenant_id or "")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    if data.get("status") == "scheduled" and not data.get("channel_jid"):
        return jsonify({"error": "Uma publicação agendada precisa do JID do canal terminado em @newsletter."}), 400
    reference = _db().collection("channel_publications").document()
    reference.set(data)
    if data.get("status") == "scheduled":
        try:
            queue_result = enqueue_publication(reference.id, data.get("scheduled_at"))
        except Exception:
            reference.set({"status": "draft", "delivery_status": "queue_unavailable", "last_error": "Fila Redis indisponível", "updated_at": _now()}, merge=True)
            return jsonify({"error": "Não foi possível agendar a publicação porque a fila está indisponível."}), 503
    else:
        queue_result = {"queued": False, "scheduled": False}
    _audit("channel_publication_created", _identity(), tenant_id, {"publication_id": reference.id, "scheduled": data.get("status") == "scheduled"})
    response_data = {"id": reference.id, **data, "queue": queue_result}
    return jsonify({"created": True, "publication": response_data, "capability": channel_capability()}), 201


@platform_bp.post("/client/channel-publications/<publication_id>/actions/<action>")
@_require_tenant_roles("owner", "operator")
def channel_publication_action(publication_id: str, action: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("channel_publications").document(publication_id)
    document = reference.get()
    data = document.to_dict() if document.exists else {}
    if not document.exists or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Publicação não encontrada."}), 404
    if action not in {"cancel", "retry"}:
        return jsonify({"error": "Acção não suportada."}), 400
    if action == "cancel":
        reference.set({"status": "cancelled", "delivery_status": "cancelled", "updated_at": _now()}, merge=True)
        try:
            import redis
            redis.from_url(os.getenv("REDIS_URL", "redis://redis:6379/1"), decode_responses=True).set(f"negobot:channel-publication:{publication_id}:control", "cancel", ex=86400)
        except Exception:
            pass
        return jsonify({"updated": True, "publication_id": publication_id, "status": "cancelled"})
    reference.set({"status": "scheduled", "delivery_status": "queued", "last_error": None, "updated_at": _now()}, merge=True)
    try:
        queue_result = enqueue_publication(publication_id)
    except Exception:
        return jsonify({"error": "Fila Redis indisponível."}), 503
    return jsonify({"updated": True, "publication_id": publication_id, "status": "scheduled", "queue": queue_result})


@platform_bp.get("/admin/channel-publications")
@_require_roles("owner", "admin")
def admin_channel_publications():
    rows = []
    for document in _db().collection("channel_publications").limit(1000).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    rows.sort(key=lambda item: str(item.get("created_at") or ""), reverse=True)
    return jsonify({"publications": rows, "capability": channel_capability()})


@platform_bp.get("/client/campaigns")
@_require_roles("client", "operator")
def list_campaigns():
    tenant_id = _tenant_for_identity(_identity())
    rows = []
    for document in _db().collection("campaigns").where("tenant_id", "==", tenant_id).limit(100).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    rows.sort(key=lambda item: item.get("created_at", ""), reverse=True)
    return jsonify({"campaigns": rows})


@platform_bp.get("/client/campaign-settings")
@_require_roles("client", "operator")
def get_campaign_settings():
    tenant_id = _tenant_for_identity(_identity())
    data = _tenant_data(tenant_id)
    stored = data.get("campaign_settings") if isinstance(data.get("campaign_settings"), dict) else {}
    return jsonify({
        "timezone": str(stored.get("timezone") or data.get("campaign_timezone") or "Africa/Maputo"),
        "silence_start": str(stored.get("silence_start") or data.get("campaign_silence_start") or "22:00"),
        "silence_end": str(stored.get("silence_end") or data.get("campaign_silence_end") or "08:00"),
        "daily_limit": int(stored.get("daily_limit") or data.get("campaign_daily_limit") or 200),
        "min_delay_seconds": int(stored.get("min_delay_seconds") or 5),
        "max_delay_seconds": int(stored.get("max_delay_seconds") or 12),
    })


@platform_bp.patch("/client/campaign-settings")
@_require_tenant_roles("owner", "operator")
def update_campaign_settings():
    tenant_id = _tenant_for_identity(_identity())
    payload = request.get_json(silent=True) or {}
    timezone_name = str(payload.get("timezone") or "Africa/Maputo").strip()[:80]
    silence_start = str(payload.get("silence_start") or "22:00").strip()
    silence_end = str(payload.get("silence_end") or "08:00").strip()
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", silence_start) or not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", silence_end):
        return jsonify({"error": "A janela de silêncio deve usar o formato HH:MM."}), 400
    try:
        daily_limit = int(payload.get("daily_limit", 200))
        min_delay_seconds = int(payload.get("min_delay_seconds", 5))
        max_delay_seconds = int(payload.get("max_delay_seconds", 12))
    except (TypeError, ValueError):
        return jsonify({"error": "Limite e atrasos devem ser números inteiros."}), 400
    if not 1 <= daily_limit <= 10000:
        return jsonify({"error": "O limite diário deve ficar entre 1 e 10.000 mensagens."}), 400
    if not 5 <= min_delay_seconds <= max_delay_seconds <= 120:
        return jsonify({"error": "Os atrasos devem respeitar 5–120 segundos e o mínimo não pode exceder o máximo."}), 400
    settings = {"timezone": timezone_name, "silence_start": silence_start, "silence_end": silence_end, "daily_limit": daily_limit, "min_delay_seconds": min_delay_seconds, "max_delay_seconds": max_delay_seconds, "updated_at": _now()}
    _db().collection("tenants").document(tenant_id).set({"campaign_settings": settings, "campaign_timezone": timezone_name, "campaign_silence_start": silence_start, "campaign_silence_end": silence_end, "campaign_daily_limit": daily_limit}, merge=True)
    _audit("campaign_settings_updated", _identity(), tenant_id, {"daily_limit": daily_limit, "silence_start": silence_start, "silence_end": silence_end})
    return jsonify({"updated": True, **settings})


@platform_bp.post("/client/contacts/import")
@_require_roles("client", "operator")
def import_contacts():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Envie um ficheiro CSV ou XLSX."}), 400
    tenant_id = _tenant_for_identity(_identity())
    rows = []
    filename = uploaded.filename.lower()
    try:
        if filename.endswith(".csv"):
            text = uploaded.read().decode("utf-8-sig")
            rows = list(csv.DictReader(io.StringIO(text)))
        elif filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.values)
            if values:
                headers = [str(item or "").strip().lower() for item in values[0]]
                rows = [dict(zip(headers, value)) for value in values[1:]]
        else:
            return jsonify({"error": "Formato não suportado. Use CSV ou XLSX."}), 400
    except Exception:
        return jsonify({"error": "Não foi possível ler o ficheiro."}), 400
    db = _db()
    entitlements = _current_entitlements(tenant_id)
    contact_limit = int(entitlements.get("contact_limit") or 0)
    existing = {re.sub(r"\D", "", str((doc.to_dict() or {}).get("phone") or "")) for doc in db.collection("contacts").where("tenant_id", "==", tenant_id).limit(min(contact_limit + 1, 5000) if contact_limit else 5000).stream()}
    available = max(0, contact_limit - len(existing)) if contact_limit else len(rows)
    if contact_limit and available <= 0:
        return jsonify({"error": f"O teu plano permite até {contact_limit} contactos. Faz upgrade para importar mais."}), 403
    batch = db.batch()
    imported = 0
    skipped = 0
    for row in rows[:min(5000, available)]:
        normalized = {str(key).strip().lower(): value for key, value in row.items()}
        name = str(normalized.get("name") or normalized.get("nome") or "").strip()
        phone = re.sub(r"\D", "", str(normalized.get("phone") or normalized.get("telefone") or normalized.get("whatsapp") or ""))
        if len(name) < 2 or len(phone) < 8 or phone in existing:
            skipped += 1
            continue
        ref = db.collection("contacts").document()
        batch.set(ref, {"tenant_id": tenant_id, "name": name, "phone": phone, "opt_in": str(normalized.get("opt_in") or normalized.get("consentimento") or "false").lower() in {"true", "1", "sim", "yes"}, "tags": [], "created_at": _now()})
        existing.add(phone)
        imported += 1
        if imported % 400 == 0:
            batch.commit()
            batch = db.batch()
    if imported % 400:
        batch.commit()
    return jsonify({"imported": imported, "skipped": skipped, "total_rows": len(rows)})


@platform_bp.post("/client/campaigns")
@_require_roles("client", "operator")
def create_campaign():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    message = str(payload.get("message") or "").strip()
    requested_channels = payload.get("channels") or ["whatsapp"]
    channels = sorted({str(channel).strip().lower() for channel in requested_channels if str(channel).strip()})
    supported_channels = {"whatsapp", "facebook", "instagram", "tiktok", "x", "linkedin", "telegram", "email"}
    if not channels or not set(channels).issubset(supported_channels):
        return jsonify({"error": "Escolhe pelo menos um canal suportado: WhatsApp, Facebook, Instagram, TikTok, X, LinkedIn, Telegram ou email."}), 400
    language = str(payload.get("language") or "pt-MZ").strip()[:20]
    tone = str(payload.get("tone") or "profissional").strip()[:60]
    offer = str(payload.get("offer") or "").strip()[:1000]
    scheduled_at = str(payload.get("scheduled_at") or "").strip()[:80] or None
    tenant_id = _tenant_for_identity(_identity())
    db = _db()
    tenant_data = _tenant_data(tenant_id)
    entitlements = _current_entitlements(tenant_id)
    if not entitlements.get("mass_broadcast"):
        return jsonify({"error": "O disparo em massa está disponível apenas durante o trial Premium ou no plano Premium activo."}), 403
    authorized_instance = str(tenant_data.get("instance_name") or "").strip()
    requested_instance = str(payload.get("instance_name") or "").strip()
    if requested_instance and authorized_instance and requested_instance != authorized_instance:
        return jsonify({"error": "A instância seleccionada não pertence a este tenant."}), 403
    if not authorized_instance:
        return jsonify({"error": "Liga primeiro o WhatsApp deste tenant antes de criar uma campanha."}), 409
    instance_name = authorized_instance
    channel_limit = plan_channel_limit(tenant_data)
    if len(channels) > channel_limit:
        return jsonify({"error": f"O teu plano permite até {channel_limit} canal(is) por campanha. Faz upgrade ou adiciona o Pacote Canais+."}), 403
    campaign_limit = entitlements.get("campaigns_per_month")
    if campaign_limit and _count_tenant_campaigns_this_month(tenant_id or "") >= int(campaign_limit):
        return jsonify({"error": f"Atingiste o limite de {campaign_limit} campanhas este mês. Faz upgrade para continuar."}), 403
    template_id = str(payload.get("template_id") or "").strip()
    if template_id:
        template_document = db.collection("campaign_templates").document(template_id).get()
        template = template_document.to_dict() if template_document.exists else None
        if not template or template.get("tenant_id") != tenant_id or template.get("status", "active") != "active":
            return jsonify({"error": "Template não encontrado ou arquivado."}), 404
        message = str(template.get("body") or "").strip()
    if len(name) < 2 or not message or len(message) > 4000:
        return jsonify({"error": "Nome e mensagem são obrigatórios e a mensagem não pode exceder 4000 caracteres."}), 400
    segment_tags = sorted({str(item).strip().lower() for item in (payload.get("tags") or payload.get("segment_tags") or []) if str(item).strip()})[:20]
    requested_limit_raw = payload.get("recipient_limit") or payload.get("max_recipients") or 0
    try:
        recipient_limit = int(requested_limit_raw) if requested_limit_raw else int(entitlements.get("contact_limit") or 1000)
    except (TypeError, ValueError):
        return jsonify({"error": "O limite de destinatários deve ser um número inteiro."}), 400
    plan_contact_limit = int(entitlements.get("contact_limit") or 0)
    if recipient_limit < 1 or (plan_contact_limit and recipient_limit > plan_contact_limit):
        return jsonify({"error": f"O limite deve ficar entre 1 e {plan_contact_limit or 'o limite do plano'} contactos."}), 400
    include_contacts = bool(payload.get("include_contacts", True))
    include_conversations = bool(payload.get("include_conversations", False))
    group_jids = sorted({str(item).strip() for item in (payload.get("group_jids") or []) if str(item).strip()})[:50]
    if (include_contacts or include_conversations) and payload.get("consent_confirmed") is not True:
        return jsonify({"error": "Confirma que os contactos e conversas elegíveis deram opt-in antes de criar a campanha."}), 400
    if group_jids:
        if payload.get("group_authorization_confirmed") is not True:
            return jsonify({"error": "Confirma que autorizas o envio apenas para os teus grupos próprios."}), 400
        authorized = set(authorized_group_jids(tenant_id or "", authorized_instance))
        unauthorized = sorted(set(group_jids) - authorized)
        if unauthorized:
            return jsonify({"error": "Um ou mais grupos não estão verificados como grupos próprios administrados pela instância deste tenant.", "unauthorized_groups": unauthorized}), 403
    contacts = []
    conversation_eligible_count = 0
    eligible_contacts_by_phone = {}
    if include_contacts or include_conversations:
        contact_documents = list(db.collection("contacts").where("tenant_id", "==", tenant_id).limit(5000).stream())
        for contact in contact_documents:
            data = contact.to_dict() or {}
            phone = re.sub(r"\D", "", str(data.get("phone") or ""))
            if not phone or data.get("opt_in") is not True or data.get("do_not_contact"):
                continue
            if segment_tags and not set(segment_tags).issubset(set(data.get("tags") or [])):
                continue
            eligible_contacts_by_phone[phone] = contact
        if include_contacts:
            contacts = list(eligible_contacts_by_phone.values())[:recipient_limit]
        if include_conversations:
            conversation_documents = db.collection("clientes_bot").document(tenant_id).collection("conversas").limit(500).stream()
            for conversation_document in conversation_documents:
                phone = re.sub(r"\D", "", str(conversation_document.id or ""))
                if phone not in eligible_contacts_by_phone:
                    continue
                conversation_eligible_count += 1
                if not include_contacts and len(contacts) < recipient_limit:
                    contacts.append(eligible_contacts_by_phone[phone])
    if not contacts and not group_jids:
        return jsonify({"error": "Adiciona contactos/conversas com opt-in ou selecciona pelo menos um grupo próprio verificado."}), 400
    campaign_ref = db.collection("campaigns").document()
    campaign_ref.set({
        "tenant_id": tenant_id,
        "name": name,
        "message": message,
        "template_id": template_id or None,
        "segment_tags": segment_tags,
        "instance_name": instance_name,
        "channels": channels,
        "language": language,
        "tone": tone,
        "offer": offer,
        "scheduled_at": scheduled_at,
        "orchestration_status": "queued",
        "status": "queued",
        "total": len(contacts),
        "sent": 0,
        "failed": 0,
        "created_at": _now(),
        "started_at": None,
        "finished_at": None,
        "include_contacts": include_contacts,
        "include_conversations": include_conversations,
        "conversation_count": conversation_eligible_count,
        "group_jids": group_jids,
        "contacts_count": 0,
        "groups_count": len(group_jids),
    })
    recipients = []
    seen_contact_ids = set()
    for contact in contacts:
        if contact.id in seen_contact_ids:
            continue
        seen_contact_ids.add(contact.id)
        data = contact.to_dict() or {}
        phone = re.sub(r"\D", "", str(data.get("phone") or ""))
        if not phone:
            continue
        recipient_ref = db.collection("campaign_recipients").document(f"{campaign_ref.id}_{contact.id}")
        recipient_ref.set({"tenant_id": tenant_id, "campaign_id": campaign_ref.id, "contact_id": contact.id, "recipient_type": "contact", "phone": phone, "status": "queued", "attempts": 0})
        recipients.append(phone)
    for group_jid in group_jids:
        safe_id = re.sub(r"[^A-Za-z0-9]+", "_", group_jid).strip("_")[:90]
        recipient_ref = db.collection("campaign_recipients").document(f"{campaign_ref.id}_group_{safe_id}")
        recipient_ref.set({"tenant_id": tenant_id, "campaign_id": campaign_ref.id, "recipient_type": "group", "group_jid": group_jid, "phone": group_jid, "group_authorized": True, "status": "queued", "attempts": 0})
        recipients.append(group_jid)
    campaign_ref.set({"total": len(recipients), "contacts_count": len([item for item in recipients if "@g.us" not in item]), "groups_count": len(group_jids)}, merge=True)
    try:
        import redis
        queue = redis.from_url(os.getenv("REDIS_URL", "redis://redis:6379/1"), decode_responses=True)
        queue.rpush("negobot:campaigns", campaign_ref.id)
    except Exception as exc:
        campaign_ref.set({"status": "failed", "error": "Fila indisponível"}, merge=True)
        return jsonify({"error": "Não foi possível iniciar a fila da campanha."}), 503
    return jsonify({"created": True, "campaign": {"id": campaign_ref.id, "name": name, "status": "queued", "total": len(recipients), "channels": channels, "scheduled_at": scheduled_at, "contacts_count": len(contacts), "groups_count": len(group_jids)}}), 201


@platform_bp.post("/client/campaigns/<campaign_id>/actions/<action>")
@_require_roles("client", "operator")
def campaign_action(campaign_id: str, action: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("campaigns").document(campaign_id)
    document = reference.get()
    if not document.exists or (document.to_dict() or {}).get("tenant_id") != tenant_id:
        return jsonify({"error": "Campanha não encontrada."}), 404
    if action not in {"pause", "resume", "cancel"}:
        return jsonify({"error": "Ação não suportada."}), 400
    state = {"pause": "paused", "resume": "queued", "cancel": "cancelled"}[action]
    reference.set({"status": state, "updated_at": _now()}, merge=True)
    try:
        import redis
        queue = redis.from_url(os.getenv("REDIS_URL", "redis://redis:6379/1"), decode_responses=True)
        queue.set(f"negobot:campaign:{campaign_id}:control", action, ex=86400)
        if action == "resume":
            queue.rpush("negobot:campaigns", campaign_id)
    except Exception:
        pass
    return jsonify({"updated": True, "campaign_id": campaign_id, "status": state})


_INTEGRATION_DEFAULTS = {
    "evolution": {"label": "Evolution API", "kind": "WhatsApp e webhooks", "env": "EVOLUTION_API_KEY"},
    "ai_worker": {"label": "AI Worker", "kind": "Conversação, transcrição e geração via fila", "env": "REDIS_URL"},
    "firebase": {"label": "Firebase Firestore", "kind": "Dados e histórico", "env": "FIREBASE_CONFIG"},
    "redis": {"label": "Redis Queue", "kind": "Fila de campanhas", "env": "REDIS_URL"},
    "n8n": {"label": "n8n Automations", "kind": "Workflows e automações", "env": "N8N_HOST"},
    "lemonsqueezy": {"label": "Lemon Squeezy", "kind": "Checkout online e subscrições", "env": "LEMONSQUEEZY_API_KEY"},
}


@platform_bp.get("/admin/integrations")
@_require_roles("owner", "admin")
def list_integrations():
    db = _db()
    stored = {doc.id: (doc.to_dict() or {}) for doc in db.collection("platform_integrations").stream()}
    rows = []
    for key, default in _INTEGRATION_DEFAULTS.items():
        item = {"key": key, **default, **stored.get(key, {})}
        configured = True if key in {"redis", "ai_worker"} else bool(os.getenv(default["env"], "").strip())
        item["configured"] = configured
        item.pop("env", None)
        rows.append(item)
    return jsonify({"integrations": rows})


@platform_bp.patch("/admin/integrations/<integration_key>")
@_require_roles("owner", "admin")
def update_integration(integration_key: str):
    if integration_key not in _INTEGRATION_DEFAULTS:
        return jsonify({"error": "Integração não encontrada."}), 404
    payload = request.get_json(silent=True) or {}
    allowed = {key: str(payload[key]).strip() for key in ("label", "public_url", "notes") if key in payload}
    if "label" in allowed and not 2 <= len(allowed["label"]) <= 80:
        return jsonify({"error": "O nome da integração deve ter entre 2 e 80 caracteres."}), 400
    _db().collection("platform_integrations").document(integration_key).set({**allowed, "updated_at": _now(), "updated_by": _identity().get("id")}, merge=True)
    return jsonify({"updated": True, "key": integration_key, "fields": sorted(allowed)})


@platform_bp.get("/client/integration/status")
@_require_roles("client", "operator")
def client_integration_status():
    tenant_id = _tenant_for_identity(_identity())
    tenant_doc = _db().collection("tenants").document(tenant_id).get()
    tenant = tenant_doc.to_dict() or {}
    instance_name = tenant.get("instance_name") or tenant_id
    base_url = str(os.getenv("EVOLUTION_API_URL", "")).rstrip("/")
    api_key = os.getenv("EVOLUTION_API_KEY", "")
    state = "not_configured"
    if base_url and api_key:
        try:
            response = requests.get(f"{base_url}/instance/connectionState/{instance_name}", headers={"apikey": api_key}, timeout=8)
            if response.ok:
                state = str((response.json() or {}).get("instance", {}).get("state") or (response.json() or {}).get("state") or "unknown")
            else:
                state = "offline"
        except requests.RequestException:
            state = "offline"
    services: dict[str, dict[str, str]] = {
        "whatsapp": {"label": "WhatsApp / Evolution", "status": "online" if state == "open" else state},
        "redis": {"label": "Campaign queue", "status": "unknown"},
        "automation": {"label": "Campaign automation", "status": "unknown"},
        "payments_local": {"label": "M-Pesa / AutoPay", "status": "configured" if os.getenv("MPESA_RECEIVER_PHONE", "855000929").strip() else "not_configured"},
        "payments_international": {"label": "Lemon Squeezy", "status": "configured" if os.getenv("LEMONSQUEEZY_API_KEY", "").strip() and os.getenv("LEMONSQUEEZY_WEBHOOK_SECRET", "").strip() else "not_configured"},
    }
    worker_profiles = ("ai", "image", "audio", "social", "mailer", "video", "campaign")
    try:
        import redis
        queue = redis.from_url(os.getenv("REDIS_URL", "redis://redis:6379/1"), decode_responses=True)
        queue.ping()
        services["redis"]["status"] = "online"
        now = time.time()
        n8n_flag = queue.get("negobot:worker:campaign:n8n_configured")
        if n8n_flag is not None:
            services["automation"]["status"] = "configured" if n8n_flag == "1" else "not_configured"
        for profile in worker_profiles:
            raw_heartbeat = queue.get(f"negobot:worker:heartbeat:{profile}")
            if raw_heartbeat:
                try:
                    services[f"worker_{profile}"] = {"label": f"{profile.title()} worker", "status": "online" if now - float(raw_heartbeat) <= 120 else "offline"}
                except (TypeError, ValueError):
                    services[f"worker_{profile}"] = {"label": f"{profile.title()} worker", "status": "unknown"}
            else:
                services[f"worker_{profile}"] = {"label": f"{profile.title()} worker", "status": "offline"}
    except Exception:
        services["redis"]["status"] = "offline"
        for profile in worker_profiles:
            services[f"worker_{profile}"] = {"label": f"{profile.title()} worker", "status": "unknown"}
    return jsonify({"instance_name": instance_name, "state": state, "configured": bool(tenant.get("instance_name")), "services": services})


@platform_bp.get("/client/assistant")
@_require_roles("client", "operator")
def get_assistant_settings():
    tenant_id = _tenant_for_identity(_identity())
    document = _db().collection("tenants").document(tenant_id).get()
    data = document.to_dict() or {}
    return jsonify({
        "diretrizes_corporativas": data.get("diretrizes_corporativas", ""),
        "base_conhecimento_documentos": data.get("base_conhecimento_documentos", ""),
        "timeout_humano_minutos": data.get("timeout_humano_minutos", 15),
        "models": {"text": "AI Worker", "vision": "Image Worker"},
    })


@platform_bp.patch("/client/assistant")
@_require_roles("client", "operator")
def update_assistant_settings():
    tenant_id = _tenant_for_identity(_identity())
    payload = request.get_json(silent=True) or {}
    allowed = {}
    if "diretrizes_corporativas" in payload:
        allowed["diretrizes_corporativas"] = str(payload["diretrizes_corporativas"])[:12000]
    if "base_conhecimento_documentos" in payload:
        allowed["base_conhecimento_documentos"] = str(payload["base_conhecimento_documentos"])[:12000]
    if "timeout_humano_minutos" in payload:
        try:
            allowed["timeout_humano_minutos"] = max(1, min(240, int(payload["timeout_humano_minutos"])))
        except (TypeError, ValueError):
            return jsonify({"error": "Timeout inválido."}), 400
    if not allowed:
        return jsonify({"error": "Nenhuma definição válida foi enviada."}), 400
    _db().collection("tenants").document(tenant_id).set({**allowed, "updated_at": _now()}, merge=True)
    return jsonify({"updated": True, "fields": sorted(allowed)})


@platform_bp.get("/client/assistant/knowledge")
@_require_roles("client", "operator")
def list_assistant_knowledge():
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "tenant não configurado"}), 403
    try:
        files = list_tenant_files(_db(), tenant_id)
    except Exception:
        logger.exception("Falha ao listar Base de Conhecimento tenant=%s", tenant_id)
        return jsonify({"error": "Não foi possível carregar a base de conhecimento neste momento."}), 503
    return jsonify({"files": files, "count": len(files)})


@platform_bp.post("/client/assistant/knowledge")
@_require_tenant_roles("owner", "operator")
def upload_assistant_knowledge():
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "tenant não configurado"}), 403
    uploaded = request.files.get("file")
    if uploaded is None or not uploaded.filename:
        return jsonify({"error": "Selecciona um ficheiro para a base de conhecimento."}), 400

    content = uploaded.read()
    try:
        filename, extension = validate_upload(uploaded.filename, uploaded.mimetype, len(content))
    except KnowledgeBaseError as exc:
        return jsonify({"error": str(exc)}), 400

    file_id = secrets.token_urlsafe(16)
    now = _now()
    db = _db()
    file_ref = db.collection("assistant_knowledge_files").document(file_id)
    base_data = {
        "tenant_id": tenant_id,
        "file_name": filename,
        "extension": extension,
        "mime_type": str(uploaded.mimetype or "application/octet-stream").split(";", 1)[0].lower(),
        "size_bytes": len(content),
        "status": "processing",
        "created_at": now,
        "updated_at": now,
    }
    file_ref.set(base_data, merge=True)

    storage_key = store_original(tenant_id, file_id, filename, content, base_data["mime_type"])
    if not storage_key:
        file_ref.set({"status": "error", "error": "Não foi possível armazenar o ficheiro.", "updated_at": _now()}, merge=True)
        return jsonify({"error": "Não foi possível armazenar o ficheiro."}), 502

    try:
        extracted = extract_text(filename, content)
    except KnowledgeBaseError as exc:
        delete_original(storage_key)
        file_ref.set({"storage_key": storage_key, "status": "error", "error": str(exc), "updated_at": _now()}, merge=True)
        return jsonify({"error": str(exc), "file": serialise_file(file_id, {**base_data, "storage_key": storage_key, "status": "error", "error": str(exc)})}), 422
    except Exception:
        logger.exception("Falha inesperada a processar ficheiro de conhecimento tenant=%s", tenant_id)
        delete_original(storage_key)
        error = "Não foi possível processar este ficheiro. Tenta novamente com um ficheiro válido."
        file_ref.set({"storage_key": storage_key, "status": "error", "error": error, "updated_at": _now()}, merge=True)
        return jsonify({"error": error}), 422

    indexed_at = _now()
    file_ref.set({
        "storage_key": storage_key,
        "extracted_text": extracted,
        "extracted_chars": len(extracted),
        "status": "indexed",
        "error": None,
        "indexed_at": indexed_at,
        "updated_at": indexed_at,
    }, merge=True)
    data = {**base_data, "storage_key": storage_key, "extracted_chars": len(extracted), "status": "indexed", "indexed_at": indexed_at}
    return jsonify({"uploaded": True, "file": serialise_file(file_id, data)}), 201


@platform_bp.delete("/client/assistant/knowledge/<file_id>")
@_require_tenant_roles("owner", "operator")
def delete_assistant_knowledge(file_id: str):
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "tenant não configurado"}), 403
    clean_file_id = str(file_id or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{8,64}", clean_file_id):
        return jsonify({"error": "Ficheiro inválido."}), 400
    ref = _db().collection("assistant_knowledge_files").document(clean_file_id)
    snapshot = ref.get()
    data = snapshot.to_dict() if snapshot.exists else {}
    if not snapshot.exists or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Ficheiro não encontrado neste tenant."}), 404
    delete_original(data.get("storage_key"))
    ref.delete()
    return jsonify({"deleted": True, "file_id": clean_file_id})


@platform_bp.get("/client/campaign-audience/conversations")
@_require_roles("client", "operator")
def campaign_conversation_audience():
    """Return only existing conversations that can be used as an opted-in campaign audience."""
    tenant_id = _tenant_for_identity(_identity())
    db = _db()
    contacts_by_phone = {}
    for document in db.collection("contacts").where("tenant_id", "==", tenant_id).limit(5000).stream():
        data = document.to_dict() or {}
        phone = re.sub(r"\D", "", str(data.get("phone") or ""))
        if phone and data.get("opt_in") is True and not data.get("do_not_contact"):
            contacts_by_phone[phone] = {"id": document.id, "name": str(data.get("name") or "").strip(), "tags": data.get("tags") or []}
    rows = []
    seen = set()
    for document in db.collection("clientes_bot").document(tenant_id).collection("conversas").limit(500).stream():
        phone = re.sub(r"\D", "", str(document.id or ""))
        contact = contacts_by_phone.get(phone)
        if not phone or not contact or phone in seen:
            continue
        seen.add(phone)
        data = document.to_dict() or {}
        last_interaction = data.get("ultima_interacao") or data.get("updated_at")
        rows.append({
            "id": phone,
            "phone": phone,
            "name": str(data.get("name") or contact.get("name") or "Contacto").strip(),
            "last_message": str(data.get("ultima_mensagem") or data.get("last_message") or "").strip(),
            "last_interaction": str(last_interaction) if last_interaction else None,
            "status_atendimento": data.get("status_atendimento") or "bot",
            "contact_id": contact["id"],
        })
    rows.sort(key=lambda item: item.get("last_interaction") or "", reverse=True)
    return jsonify({"conversations": rows, "count": len(rows), "eligibility": "opt_in_contact_only"})


def _conversation_sources(tenant_id: str, instance_name: str) -> list[Any]:
    sources = [_db().collection("clientes_bot").document(tenant_id).collection("conversas")]
    if instance_name and instance_name != tenant_id:
        sources.append(_db().collection("clientes_bot").document(instance_name).collection("conversas"))
    return sources


def _clean_chat_target(value: str) -> str:
    raw = str(value or "").strip()
    lowered = raw.lower()
    if lowered.endswith("@g.us"):
        return raw
    if lowered.endswith("@lid"):
        return ""
    if "@" in raw and not lowered.endswith("@s.whatsapp.net"):
        return ""
    return re.sub(r"\D", "", raw)


def _chat_target_from_evolution(chat: dict[str, Any]) -> str:
    """Extract a complete WhatsApp JID/number, never a short @lid or display fragment."""
    candidates: list[tuple[str, Any]] = [
        ("remoteJid", chat.get("remoteJid")),
        ("jid", chat.get("jid")),
        ("id", chat.get("id")),
        ("phone", chat.get("phone")),
        ("number", chat.get("number")),
        ("phoneNumber", chat.get("phoneNumber")),
    ]
    for nested in (
        chat.get("key"),
        chat.get("lastMessage"),
        chat.get("contact"),
        chat.get("profile"),
    ):
        if isinstance(nested, dict):
            candidates.extend([
                ("remoteJid", nested.get("remoteJid")),
                ("jid", nested.get("jid")),
                ("id", nested.get("id")),
                ("phone", nested.get("phone")),
                ("number", nested.get("number")),
                ("phoneNumber", nested.get("phoneNumber")),
                ("remoteJid", (nested.get("key") or {}).get("remoteJid") if isinstance(nested.get("key"), dict) else None),
            ])
    for field, candidate in candidates:
        raw = str(candidate or "").strip()
        lowered = raw.lower()
        if lowered.endswith("@g.us"):
            return raw
        if lowered.endswith("@lid") or ("@" in raw and not lowered.endswith("@s.whatsapp.net")):
            continue
        cleaned = _clean_chat_target(raw)
        if len(cleaned) >= 8:
            return cleaned
    return ""


def _real_chat_name(*values: Any) -> str:
    """Return a provider/user name, ignoring placeholders created by old sync code."""
    placeholders = {"contacto", "contact", "cliente", "customer", "unknown", "undefined", "null", "sem nome"}
    for value in values:
        name = str(value or "").strip()
        if name and name.casefold() not in placeholders:
            return name
    return ""


def _chat_display_name(chat: dict[str, Any]) -> str:
    contact = chat.get("contact") if isinstance(chat.get("contact"), dict) else {}
    profile = chat.get("profile") if isinstance(chat.get("profile"), dict) else {}
    return _real_chat_name(
        chat.get("name"),
        chat.get("pushName"),
        chat.get("subject"),
        chat.get("displayName"),
        contact.get("name"),
        contact.get("pushName"),
        profile.get("name"),
        profile.get("pushName"),
    )


def _chat_picture_url(chat: dict[str, Any]) -> str:
    for key in ("profilePictureUrl", "profilePicUrl", "picture", "pictureUrl", "avatar_url"):
        value = chat.get(key)
        if isinstance(value, str) and value.strip().startswith(("https://", "http://")):
            return value.strip()
    for nested_key in ("contact", "profile"):
        nested = chat.get(nested_key)
        if isinstance(nested, dict):
            value = _chat_picture_url(nested)
            if value:
                return value
    return ""


def _authorize_chat_target(tenant_id: str, instance_name: str, target: str, *, require_group_admin: bool = True, allow_live_instance: bool = False) -> tuple[bool, int, str]:
    if target.endswith("@g.us"):
        group_doc = _db().collection("whatsapp_groups").document(group_document_id(target)).get()
        group_data = group_doc.to_dict() if group_doc.exists else {}
        allowed = bool(group_data and group_data.get("tenant_id") == tenant_id)
        if require_group_admin:
            allowed = allowed and group_data.get("admin_verified") is True and group_data.get("status") == "active"
        if not allowed:
            return False, 403, "Este grupo não está autorizado para este tenant."
        return True, 200, ""
    contact_exists = target in _tenant_chat_contacts(tenant_id)
    conversation_exists = any(source.document(target).get().exists for source in _conversation_sources(tenant_id, instance_name))
    if not contact_exists and not conversation_exists and allow_live_instance and instance_name and get_connection_state(instance_name) == "open":
        live_targets = {_chat_target_from_evolution(chat) for chat in listar_chats_whatsapp(instance_name)}
        if target in live_targets:
            return True, 200, ""
    if not contact_exists and not conversation_exists:
        return False, 403, "Só podes consultar uma conversa pertencente a este tenant."
    return True, 200, ""


def _serialize_chat_message(document: Any, data: dict[str, Any]) -> dict[str, Any]:
    stamp = data.get("timestamp") or data.get("created_at") or data.get("updated_at")
    if hasattr(stamp, "isoformat"):
        stamp = stamp.isoformat()
    text = str(data.get("text") or data.get("content") or data.get("caption") or "").strip()
    role = str(data.get("role") or "user").strip().lower()
    raw_media_type = str(data.get("media_type") or data.get("mediatype") or data.get("message_type") or "").strip().lower()
    media_type = "image" if raw_media_type in {"image", "imagemessage", "image_message", "photo"} else "document" if raw_media_type in {"document", "documentmessage", "document_message", "file"} else (raw_media_type or None)
    raw_media_url = data.get("media_url") or data.get("url")
    media_url = raw_media_url.strip() if isinstance(raw_media_url, str) and raw_media_url.strip().startswith(("https://", "http://")) else None
    return {
        "id": getattr(document, "id", None) or str(data.get("id") or ""),
        "role": role,
        "text": text,
        "timestamp": str(stamp) if stamp else None,
        "from_me": role in {"assistant", "atendente", "bot", "agent"},
        "media_type": media_type,
        "media_url": media_url,
        "file_name": str(data.get("file_name") or data.get("filename") or data.get("fileName") or "").strip() or None,
        "mime_type": str(data.get("mime_type") or data.get("mimetype") or "").strip().lower() or None,
        "caption": str(data.get("caption") or "").strip() or None,
    }


def _tenant_chat_groups(tenant_id: str) -> dict[str, dict[str, Any]]:
    """Return tenant-owned group metadata indexed by both full JID and numeric prefix."""
    groups: dict[str, dict[str, Any]] = {}
    for group_doc in _db().collection("whatsapp_groups").where("tenant_id", "==", tenant_id).limit(1000).stream():
        data = group_doc.to_dict() or {}
        group_jid = str(data.get("group_jid") or "").strip()
        if not group_jid.endswith("@g.us"):
            continue
        item = {
            "id": group_doc.id,
            "group_jid": group_jid,
            "name": _real_chat_name(data.get("name"), data.get("subject"), data.get("group_name")) or group_jid,
            "admin_verified": data.get("admin_verified") is True,
            "status": data.get("status") or "unknown",
        }
        groups[group_jid] = item
        groups[re.sub(r"\D", "", group_jid.split("@", 1)[0])] = item
    return groups


def _tenant_chat_contacts(tenant_id: str) -> dict[str, dict[str, Any]]:
    """Merge platform contacts with WhatsApp contacts synced under the tenant document."""
    db = _db()
    group_phones = set(_tenant_chat_groups(tenant_id))
    contacts: dict[str, dict[str, Any]] = {}
    for contact_doc in db.collection("contacts").where("tenant_id", "==", tenant_id).limit(5000).stream():
        data = contact_doc.to_dict() or {}
        normalized = _clean_chat_target(str(data.get("phone") or data.get("telefone") or ""))
        if normalized and normalized not in group_phones and not normalized.endswith("@g.us") and len(normalized) >= 8:
            contacts[normalized] = {
                "id": contact_doc.id,
                "name": _real_chat_name(data.get("name"), data.get("nome"), data.get("pushName")),
            }
    base_contacts = db.collection("clientes_bot").document(tenant_id).collection("base_contactos").limit(5000).stream()
    for contact_doc in base_contacts:
        data = contact_doc.to_dict() or {}
        normalized = _clean_chat_target(str(data.get("phone") or data.get("telefone") or contact_doc.id or ""))
        if not normalized or normalized in group_phones or normalized.endswith("@g.us") or len(normalized) < 8:
            continue
        name = _real_chat_name(data.get("name"), data.get("nome"), data.get("pushName"), data.get("display_name"))
        current = contacts.get(normalized)
        if current is None or not current.get("name"):
            contacts[normalized] = {"id": contact_doc.id, "name": name}
    return contacts


@platform_bp.get("/client/conversations")
@_require_roles("client", "operator")
def list_conversations():
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    tenant_groups = _tenant_chat_groups(tenant_id)
    contacts = _tenant_chat_contacts(tenant_id)
    by_phone: dict[str, dict[str, Any]] = {}
    for phone, contact in contacts.items():
        if phone in tenant_groups:
            continue
        by_phone[phone] = {
            "id": phone,
            "phone": phone,
            "name": _real_chat_name(contact.get("name")) or phone,
            "last_message": "",
            "last_interaction": None,
            "status_atendimento": "bot",
            "contact_id": contact.get("id"),
            "kind": "contact",
        }
    if instance_name and get_connection_state(instance_name) == "open":
        for chat in listar_chats_whatsapp(instance_name):
            phone = _chat_target_from_evolution(chat)
            if not phone:
                continue
            numeric_phone = re.sub(r"\D", "", phone.split("@", 1)[0]) if phone.endswith("@g.us") else phone
            group = tenant_groups.get(phone) or tenant_groups.get(numeric_phone)
            if group:
                phone = group["group_jid"]
            last_message_data = chat.get("lastMessage") or chat.get("last_message") or {}
            if not isinstance(last_message_data, dict):
                last_message_data = {}
            last_message = str(
                chat.get("lastMessageText")
                or chat.get("last_message_text")
                or last_message_data.get("conversation")
                or last_message_data.get("text")
                or ((last_message_data.get("extendedTextMessage") or {}).get("text") if isinstance(last_message_data.get("extendedTextMessage"), dict) else "")
                or ""
            ).strip()
            last_interaction = chat.get("updatedAt") or chat.get("timestamp") or chat.get("conversationTimestamp")
            if hasattr(last_interaction, "isoformat"):
                last_interaction = last_interaction.isoformat()
            by_phone.setdefault(phone, {
                "id": phone,
                "phone": phone,
                "name": _chat_display_name(chat) or _real_chat_name((group or {}).get("name"), contacts.get(phone, {}).get("name")) or phone,
                "last_message": last_message,
                "last_interaction": str(last_interaction) if last_interaction else None,
                "status_atendimento": "bot",
                "contact_id": contacts.get(phone, {}).get("id"),
                "avatar_url": _chat_picture_url(chat) or None,
                "kind": "group" if phone.endswith("@g.us") else "contact",
            })
    for source in _conversation_sources(tenant_id, instance_name):
        for document in source.limit(5000).stream():
            data = document.to_dict() or {}
            phone = _clean_chat_target(document.id)
            if not phone or (not phone.endswith("@g.us") and len(phone) < 8):
                continue
            numeric_phone = re.sub(r"\D", "", phone.split("@", 1)[0]) if phone.endswith("@g.us") else phone
            group = tenant_groups.get(phone) or tenant_groups.get(numeric_phone)
            if group:
                phone = group["group_jid"]
            existing = by_phone.get(phone, {})
            last_message = str(data.get("ultima_mensagem") or data.get("last_message") or existing.get("last_message") or "").strip()
            last_interaction = data.get("ultima_interacao") or data.get("updated_at") or existing.get("last_interaction")
            if hasattr(last_interaction, "isoformat"):
                last_interaction = last_interaction.isoformat()
            by_phone[phone] = {
                **existing,
                "id": phone,
                "phone": phone,
                "name": _real_chat_name(data.get("name"), (group or {}).get("name"), contacts.get(phone, {}).get("name"), existing.get("name")) or phone,
                "last_message": last_message,
                "last_interaction": str(last_interaction) if last_interaction else None,
                "status_atendimento": data.get("status_atendimento") or existing.get("status_atendimento") or "bot",
                "contact_id": contacts.get(phone, {}).get("id"),
                "kind": "group" if phone.endswith("@g.us") else "contact",
            }
    rows = sorted(by_phone.values(), key=lambda item: str(item.get("last_interaction") or ""), reverse=True)
    return jsonify({"conversations": rows[:500], "count": len(rows), "instance_name": instance_name})


@platform_bp.get("/client/conversations/<phone>/profile")
@_require_roles("client", "operator")
def conversation_profile(phone: str):
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    target = _clean_chat_target(phone)
    if not target or (not target.endswith("@g.us") and len(target) < 8):
        return jsonify({"error": "Destino de conversa inválido."}), 400
    allowed, status_code, error = _authorize_chat_target(tenant_id, instance_name, target, require_group_admin=True, allow_live_instance=True)
    if not allowed:
        return jsonify({"error": error}), status_code
    if not instance_name or get_connection_state(instance_name) != "open":
        return jsonify({"phone": target, "profile_picture_url": None, "available": False})
    remote_url = get_profile_picture_url(target, instance_name=instance_name)
    proxy_url = f"/api/platform/client/conversations/{quote(target, safe='')}/profile/image" if remote_url else None
    return jsonify({"phone": target, "profile_picture_url": proxy_url, "available": bool(remote_url)})


@platform_bp.get("/client/conversations/<phone>/profile/image")
@_require_roles("client", "operator")
def conversation_profile_image(phone: str):
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    target = _clean_chat_target(phone)
    if not target or (not target.endswith("@g.us") and len(target) < 8):
        return jsonify({"error": "Destino de conversa inválido."}), 400
    allowed, status_code, error = _authorize_chat_target(tenant_id, instance_name, target, require_group_admin=True, allow_live_instance=True)
    if not allowed:
        return jsonify({"error": error}), status_code
    if not instance_name or get_connection_state(instance_name) != "open":
        return jsonify({"error": "O WhatsApp deste tenant está desligado."}), 409
    remote_url = get_profile_picture_url(target, instance_name=instance_name)
    if not remote_url:
        return jsonify({"error": "Este contacto não tem uma foto de perfil disponível."}), 404
    try:
        image_response = requests.get(remote_url, timeout=15)
        image_response.raise_for_status()
        content_type = image_response.headers.get("Content-Type", "image/jpeg").split(";", 1)[0]
        if not content_type.startswith("image/"):
            return jsonify({"error": "A Evolution devolveu um conteúdo que não é uma imagem."}), 502
        response = Response(image_response.content, mimetype=content_type)
        response.headers["Cache-Control"] = "private, max-age=300"
        return response
    except requests.RequestException:
        return jsonify({"error": "Não foi possível carregar a foto de perfil."}), 502


@platform_bp.get("/client/conversations/<phone>/messages")
@_require_roles("client", "operator")
def conversation_messages(phone: str):
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    target = _clean_chat_target(phone)
    if not target or (not target.endswith("@g.us") and len(target) < 8):
        return jsonify({"error": "Destino de conversa inválido."}), 400
    allowed, status_code, error = _authorize_chat_target(tenant_id, instance_name, target, require_group_admin=True, allow_live_instance=True)
    if not allowed:
        return jsonify({"error": error}), status_code
    messages: list[dict[str, Any]] = []
    seen: set[str] = set()
    for source in _conversation_sources(tenant_id, instance_name):
        conversation_ref = source.document(target)
        conversation_doc = conversation_ref.get()
        if conversation_doc.exists:
            data = conversation_doc.to_dict() or {}
            inline = data.get("messages") if isinstance(data.get("messages"), list) else []
            for index, item in enumerate(inline):
                if isinstance(item, dict):
                    serialized = _serialize_chat_message(type("Message", (), {"id": f"inline-{index}"})(), item)
                    key = f"{serialized.get('timestamp')}|{serialized.get('text')}|{serialized.get('role')}"
                    if key not in seen:
                        seen.add(key); messages.append(serialized)
        for document in conversation_ref.collection("historico").limit(200).stream():
            serialized = _serialize_chat_message(document, document.to_dict() or {})
            key = f"{serialized.get('timestamp')}|{serialized.get('text')}|{serialized.get('role')}"
            if serialized.get("text") and key not in seen:
                seen.add(key); messages.append(serialized)
    messages.sort(key=lambda item: str(item.get("timestamp") or ""))
    return jsonify({"phone": target, "messages": messages[-200:], "count": min(len(messages), 200)})


@platform_bp.post("/client/conversations/<phone>/media")
@_require_tenant_roles("owner", "operator")
def send_conversation_media(phone: str):
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    target = _clean_chat_target(phone)
    if not target or (not target.endswith("@g.us") and len(target) < 8):
        return jsonify({"error": "Destino de conversa inválido."}), 400
    if not instance_name:
        return jsonify({"error": "Liga primeiro o WhatsApp deste tenant."}), 409
    if get_connection_state(instance_name) != "open":
        return jsonify({"error": "O WhatsApp deste tenant está desligado. Liga a instância antes de enviar ficheiros."}), 409
    allowed, status_code, error = _authorize_chat_target(tenant_id, instance_name, target, require_group_admin=True, allow_live_instance=True)
    if not allowed:
        return jsonify({"error": error}), status_code
    uploaded = request.files.get("file")
    if uploaded is None or not uploaded.filename:
        return jsonify({"error": "Selecciona uma imagem ou documento."}), 400
    mime_type = str(uploaded.mimetype or "application/octet-stream").lower().strip()
    allowed_image = mime_type.startswith("image/")
    allowed_document = mime_type in {
        "application/pdf", "application/msword", "application/rtf", "application/zip",
        "application/vnd.ms-excel", "application/vnd.ms-powerpoint",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "text/plain", "text/csv",
    }
    if not allowed_image and not allowed_document:
        return jsonify({"error": "Apenas imagens e documentos PDF, Word, Excel, PowerPoint, CSV ou TXT são permitidos."}), 415
    content = uploaded.stream.read(16 * 1024 * 1024 + 1)
    if len(content) > 16 * 1024 * 1024:
        return jsonify({"error": "O ficheiro não pode exceder 16 MB."}), 413
    filename = secure_filename(uploaded.filename)[:180] or ("imagem" if allowed_image else "documento")
    caption = str(request.form.get("caption") or "").strip()[:4000]
    media_type = "image" if allowed_image else "document"
    encoded = f"data:{mime_type};base64,{base64.b64encode(content).decode('ascii')}"
    if not send_media(target, encoded, caption=caption, mediatype=media_type, filename=filename, mimetype=mime_type, instance_name=instance_name):
        return jsonify({"error": "A Evolution API não conseguiu enviar o ficheiro."}), 502
    now = _now()
    label = caption or ("Imagem enviada" if allowed_image else f"Documento enviado: {filename}")
    conversation_ref = _db().collection("clientes_bot").document(instance_name).collection("conversas").document(target)
    conversation_ref.set({"ultima_mensagem": label, "ultima_mensagem_por": "atendente", "ultima_interacao": now, "status_atendimento": "humano"}, merge=True)
    message_data = {"role": "atendente", "text": caption, "caption": caption, "media_type": media_type, "file_name": filename, "mime_type": mime_type, "timestamp": now}
    conversation_ref.collection("historico").add(message_data)
    return jsonify({"sent": True, "phone": target, "message": _serialize_chat_message(type("Message", (), {"id": "media-outgoing"})(), message_data)})


@platform_bp.post("/client/conversations/<phone>/messages")
@_require_tenant_roles("owner", "operator")
def send_conversation_message(phone: str):
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    instance_name = str(tenant.get("instance_name") or "").strip()
    target = _clean_chat_target(phone)
    payload = request.get_json(silent=True) or {}
    text = str(payload.get("text") or "").strip()
    if not target or (not target.endswith("@g.us") and len(target) < 8):
        return jsonify({"error": "Destino de conversa inválido."}), 400
    if not text or len(text) > 4000:
        return jsonify({"error": "A mensagem deve ter entre 1 e 4.000 caracteres."}), 400
    if not instance_name:
        return jsonify({"error": "Liga primeiro o WhatsApp deste tenant."}), 409
    if get_connection_state(instance_name) != "open":
        return jsonify({"error": "O WhatsApp deste tenant está desligado. Liga a instância antes de enviar mensagens."}), 409
    if target.endswith("@g.us"):
        group_doc = _db().collection("whatsapp_groups").document(group_document_id(target)).get()
        group_data = group_doc.to_dict() if group_doc.exists else {}
        if not group_data or group_data.get("tenant_id") != tenant_id or group_data.get("admin_verified") is not True or group_data.get("status") != "active":
            return jsonify({"error": "Este grupo não está verificado como grupo próprio administrado pela instância deste tenant."}), 403
    allowed, status_code, error = _authorize_chat_target(tenant_id, instance_name, target, require_group_admin=True, allow_live_instance=True)
    if not allowed:
        return jsonify({"error": error.replace("consultar", "iniciar")}), status_code
    if not send_whatsapp(target, text, instance_name=instance_name):
        return jsonify({"error": "A Evolution API não conseguiu enviar a mensagem."}), 502
    now = _now()
    conversation_ref = _db().collection("clientes_bot").document(instance_name).collection("conversas").document(target)
    conversation_ref.set({"ultima_mensagem": text, "ultima_mensagem_por": "atendente", "ultima_interacao": now, "status_atendimento": "humano"}, merge=True)
    conversation_ref.collection("historico").add({"role": "atendente", "text": text, "timestamp": now})
    return jsonify({"sent": True, "message": {"role": "atendente", "text": text, "timestamp": now.isoformat(), "from_me": True}, "phone": target})


@platform_bp.post("/client/conversations/<phone>/handoff")
@_require_roles("client", "operator")
def conversation_handoff(phone: str):
    tenant_id = _tenant_for_identity(_identity())
    clean_phone = re.sub(r"\D", "", phone)
    if len(clean_phone) < 8:
        return jsonify({"error": "Telefone inválido."}), 400
    payload = request.get_json(silent=True) or {}
    mode = str(payload.get("mode") or "").lower()
    if mode not in {"bot", "humano"}:
        return jsonify({"error": "O modo deve ser bot ou humano."}), 400
    reference = _db().collection("clientes_bot").document(tenant_id).collection("conversas").document(clean_phone)
    if not reference.get().exists:
        return jsonify({"error": "Conversa não encontrada."}), 404
    reference.set({"status_atendimento": mode, "ultima_interacao": _now()}, merge=True)
    return jsonify({"updated": True, "phone": clean_phone, "mode": mode})


@platform_bp.get("/client/plan")
@_require_roles("client", "operator")
def client_plan():
    tenant_id = _tenant_for_identity(_identity())
    data = _tenant_data(tenant_id)
    entitlements = _current_entitlements(tenant_id)
    campaign_usage = _count_tenant_campaigns_this_month(tenant_id or "")
    contact_usage = _count_tenant_contacts(tenant_id or "", int(entitlements.get("contact_limit") or 20000))
    team_usage = _count_tenant_team(tenant_id or "")
    return jsonify({
        "plan": data.get("plano", data.get("plan", entitlements["plan_id"])),
        "plan_name": entitlements["plan_name"] if entitlements.get("trial_access") else data.get("nome_plano", entitlements["plan_name"]),
        "status": data.get("status_plano", data.get("status", "demonstracao")),
        "expires_at": data.get("data_expiracao"),
        "mass_broadcast": bool(data.get("disparo_liberado", entitlements["mass_broadcast"])),
        "limits": {**(data.get("limits") or {}), **entitlements},
        "usage": {"contacts": contact_usage, "campaigns_this_month": campaign_usage, "team_seats": team_usage},
        "trial_status": data.get("trial_status", data.get("status_plano", "demonstracao")),
        "billing_region": data.get("billing_region", "mozambique"),
        "selected_plan": data.get("selected_plan"),
        "trial_access": bool(entitlements.get("trial_access", False)),
        "trial_access_level": entitlements.get("trial_access_level", "standard"),
        "trial_features": ["vídeo", "PDFs e documentos", "áudio", "imagens", "campanhas avançadas", "funcionalidades Premium"] if entitlements.get("trial_access") else [],
    })


@platform_bp.post("/client/payments/mpesa/verify")
@_require_roles("client", "operator")
def verify_mpesa_payment():
    payload = request.get_json(silent=True) or {}
    message_text = str(payload.get("message_text") or "").strip()
    client_phone = str(payload.get("client_phone") or "").strip()
    addon_id = str(payload.get("addon_id") or "").strip().lower()
    if not message_text:
        return jsonify({"error": "Introduza o código ou SMS do M-Pesa."}), 400
    tenant_id = _tenant_for_identity(_identity())
    if addon_id:
        try:
            from services.plan_service import ADDONS
            from services.payment_service import extrair_codigo_mpesa, validar_e_ativar_extra_mpesa
            addon = ADDONS.get(addon_id)
            if not addon:
                return jsonify({"error": "Extra não encontrado."}), 400
            tenant_document = _db().collection("tenants").document(tenant_id).get()
            tenant_data = tenant_document.to_dict() or {}
            if tenant_data.get("billing_region", "mozambique") == "international":
                return jsonify({"error": "Esta conta está configurada para pagamento internacional. Usa o checkout Lemon Squeezy."}), 403
            tx_id = extrair_codigo_mpesa(message_text)
            intent_ref = _db().collection("payment_intents").document()
            intent_ref.set({
                "tenant_id": tenant_id,
                "client_phone": re.sub(r"\D", "", client_phone),
                "transaction_id": tx_id,
                "addon_id": addon_id,
                "addon_name": addon["name"],
                "purchase_type": "addon",
                "status": "pending",
                "source": "platform_verify",
                "created_at": _now(),
            })
            response = validar_e_ativar_extra_mpesa(tenant_id, client_phone, message_text, addon_id)
            confirmed = "activado com sucesso" in response.lower()
            intent_ref.set({"status": "confirmed" if confirmed else "pending_validation", "transaction_id": tx_id, "updated_at": _now()}, merge=True)
            if confirmed:
                _audit("mpesa_addon_confirmed", _identity(), tenant_id, {"addon_id": addon_id, "transaction_id": tx_id})
            return jsonify({"processed": True, "response": response, "addon_id": addon_id})
        except Exception:
            return jsonify({"error": "O serviço de pagamentos está temporariamente indisponível."}), 503
    tenant_document = _db().collection("tenants").document(tenant_id).get()
    tenant_data = tenant_document.to_dict() or {}
    if tenant_data.get("billing_region", "mozambique") == "international":
        return jsonify({"error": "Esta conta está configurada para pagamento internacional. Usa o checkout Lemon Squeezy."}), 403
    intent_ref = None
    tx_id = None
    try:
        # A intenção permite ao listener AutoPay associar uma transação nova
        # ao tenant mesmo quando o Android ainda não gravou tenant_id.
        from services.payment_service import extrair_codigo_mpesa, validar_e_ativar_pagamento_mpesa
        tx_id = extrair_codigo_mpesa(message_text)
        if client_phone:
            intent_ref = _db().collection("payment_intents").document()
            intent_ref.set({
                "tenant_id": tenant_id,
                "client_phone": re.sub(r"\D", "", client_phone),
                "transaction_id": tx_id,
                "status": "pending",
                "source": "platform_verify",
                "created_at": _now(),
            })
        response = validar_e_ativar_pagamento_mpesa(tenant_id, client_phone, message_text)
        if "PAGAMENTO CONFIRMADO" in str(response):
            paid_doc = _db().collection("clientes_bot").document(tenant_id).get()
            paid = paid_doc.to_dict() or {}
            _db().collection("tenants").document(tenant_id).set({
                "plan": paid.get("plano", paid.get("plan", "demonstracao")),
                "plano": paid.get("plano", paid.get("plan", "demonstracao")),
                "plan_name": paid.get("nome_plano", "Demonstração"),
                "nome_plano": paid.get("nome_plano", "Demonstração"),
                "status": paid.get("status_plano", "ativo"),
                "status_plano": paid.get("status_plano", "ativo"),
                "data_expiracao": paid.get("data_expiracao"),
                "mass_broadcast": bool(paid.get("disparo_liberado", False)),
                "disparo_liberado": bool(paid.get("disparo_liberado", False)),
                "limite_conversas": paid.get("limite_conversas"),
                "telefone_proprietario": client_phone or paid.get("telefone_proprietario"),
                "plan_rules_version": "2026-08-v2",
                "updated_at": _now(),
            }, merge=True)
            if intent_ref is not None:
                intent_ref.set({"status": "confirmed", "transaction_id": tx_id, "confirmed_at": _now()}, merge=True)
            _audit("mpesa_payment_confirmed", _identity(), tenant_id, {"plan": paid.get("plano"), "transaction_id": tx_id, "mass_broadcast": bool(paid.get("disparo_liberado", False))})
        elif intent_ref is not None:
            intent_ref.set({"status": "pending_validation", "transaction_id": tx_id, "updated_at": _now()}, merge=True)
    except Exception:
        return jsonify({"error": "O serviço de pagamentos está temporariamente indisponível."}), 503
    qr_payload = {"state": "not_requested", "qrcode": None, "instance_name": None}
    if "PAGAMENTO CONFIRMADO" in str(response):
        try:
            owner_phone = re.sub(r"\D", "", str(client_phone or ""))
            if owner_phone:
                from services.evolution_service import criar_e_configurar_instancia_automatica, obter_qrcode_instancia
                if criar_e_configurar_instancia_automatica(owner_phone):
                    qr_data = obter_qrcode_instancia(owner_phone)
                    base64_qr = qr_data.get("base64")
                    qr_payload = {
                        "state": qr_data.get("state", "connecting"),
                        "instance_name": qr_data.get("instance_name"),
                        "qrcode": base64_qr and (base64_qr if str(base64_qr).startswith("data:") else f"data:image/png;base64,{base64_qr}"),
                    }
        except Exception:
            qr_payload["state"] = "pending"
    return jsonify({"processed": True, "response": response, **qr_payload})


@platform_bp.get("/client/payments/history")
@_require_roles("client", "operator")
def payment_history():
    tenant_id = _tenant_for_identity(_identity())
    rows = []
    for document in _db().collection("payment_intents").where("tenant_id", "==", tenant_id).limit(100).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        item.pop("message_text", None)
        rows.append(item)
    rows.sort(key=lambda item: str(item.get("created_at", "")), reverse=True)
    return jsonify({"payments": rows})



def _audit(event: str, actor: dict[str, Any] | None, tenant_id: str | None = None, metadata: dict[str, Any] | None = None):
    try:
        _db().collection("audit_events").add({
            "event": event,
            "actor_id": (actor or {}).get("id"),
            "actor_role": (actor or {}).get("role"),
            "tenant_id": tenant_id,
            "metadata": metadata or {},
            "created_at": _now(),
        })
    except Exception:
        pass


def _lemon_plan_from_variant(variant_id: str, requested_plan: str = "") -> str | None:
    from services.payment_service import TABELA_PLANOS
    variant_id = str(variant_id or "").strip()
    requested_plan = str(requested_plan or "").strip().lower()
    for data in TABELA_PLANOS.values():
        plan_id = str(data["id"])
        configured_variant = os.getenv(f"LEMONSQUEEZY_VARIANT_{plan_id.upper()}", "").strip()
        if configured_variant and configured_variant == variant_id:
            return plan_id
    if requested_plan in {str(data["id"]) for data in TABELA_PLANOS.values()} and not variant_id:
        return requested_plan
    return None


def _apply_lemon_plan(tenant_id: str, plan_id: str, attributes: dict[str, Any], intent_ref, event: dict[str, Any]):
    from services.payment_service import TABELA_PLANOS
    plan = next((item for item in TABELA_PLANOS.values() if item["id"] == plan_id), None)
    if not plan:
        raise ValueError("Plano Lemon Squeezy não encontrado.")
    now = _now()
    expires_at = __import__("services.lemonsqueezy_service", fromlist=["expiry_for_event"]).expiry_for_event(attributes, now)
    lemon_fields = {
        "provider": "lemonsqueezy",
        "payment_provider": "lemonsqueezy",
        "plan_id": plan_id,
        "plan": plan_id,
        "plan_name": plan["nome"],
        "nome_plano": plan["nome"],
        "status": "active",
        "status_plano": "ativo",
        "mass_broadcast": bool(plan["disparo_liberado"]),
        "disparo_liberado": bool(plan["disparo_liberado"]),
        "limite_conversas": plan["limite_conversas"],
        "data_ativacao": now,
        "data_expiracao": expires_at,
        "metodo_pagamento": "Lemon Squeezy",
        "lemon_subscription_id": event["object_id"] if event["object_type"] == "subscriptions" else attributes.get("subscription_id"),
        "lemon_order_id": attributes.get("order_id"),
        "lemon_customer_id": attributes.get("customer_id"),
        "lemon_variant_id": event["variant_id"],
        "lemon_status": attributes.get("status") or "active",
        "plan_rules_version": "2026-08-v2",
        "updated_at": now,
    }
    _db().collection("tenants").document(tenant_id).set(lemon_fields, merge=True)
    _db().collection("clientes_bot").document(tenant_id).set(lemon_fields, merge=True)
    intent_ref.set({**lemon_fields, "status": "confirmed", "provider_status": "active", "confirmed_at": now, "activated_at": now}, merge=True)


@platform_bp.get("/client/payments/lemonsqueezy/status")
@_require_roles("client", "operator")
def lemonsqueezy_status():
    from services.lemonsqueezy_service import configured
    from services.payment_service import TABELA_PLANOS
    from services.plan_service import ADDONS
    plans = {data["id"]: bool(os.getenv(f"LEMONSQUEEZY_VARIANT_{str(data['id']).upper()}", "").strip()) for data in TABELA_PLANOS.values()}
    addons = {addon_id: bool(os.getenv(f"LEMONSQUEEZY_VARIANT_ADDON_{addon_id.upper()}", "").strip()) for addon_id in ADDONS}
    return jsonify({"configured": configured(), "currency": os.getenv("LEMONSQUEEZY_CURRENCY", "USD"), "plans": plans, "addons": addons})


@platform_bp.post("/client/payments/lemonsqueezy/checkout")
@_require_roles("client", "operator")
def create_lemonsqueezy_checkout():
    from services.lemonsqueezy_service import create_checkout
    payload = request.get_json(silent=True) or {}
    plan_id = str(payload.get("plan_id") or "").strip().lower()
    from services.payment_service import TABELA_PLANOS
    plan = next((item for item in TABELA_PLANOS.values() if item["id"] == plan_id), None)
    if not plan:
        return jsonify({"error": "Plano não encontrado."}), 400
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "Tenant não encontrado na sessão."}), 400
    tenant_doc = _db().collection("tenants").document(tenant_id).get()
    tenant = tenant_doc.to_dict() or {}
    if tenant.get("billing_region", "mozambique") != "international":
        return jsonify({"error": "Esta conta está configurada para M-Pesa. Usa a validação AutoPay na área de pagamentos."}), 403
    intent_ref = _db().collection("payment_intents").document()
    intent_ref.set({
        "tenant_id": tenant_id,
        "provider": "lemonsqueezy",
        "payment_provider": "lemonsqueezy",
        "plan_id": plan_id,
        "plan_name": plan["nome"],
        "status": "pending_checkout",
        "created_at": _now(),
    })
    identity = _identity() or {}
    try:
        checkout = create_checkout(
            plan_id=plan_id,
            tenant_id=tenant_id,
            payment_intent_id=intent_ref.id,
            email=str(identity.get("email") or tenant.get("email") or "").strip() or None,
            name=str(identity.get("name") or tenant.get("name") or "").strip() or None,
        )
    except (RuntimeError, ValueError, requests.RequestException) as exc:
        intent_ref.set({"status": "checkout_failed", "error": str(exc), "updated_at": _now()}, merge=True)
        return jsonify({"error": "O checkout Lemon Squeezy ainda não está disponível. Usa M-Pesa ou tenta novamente mais tarde."}), 503
    intent_ref.set({"status": "pending", "checkout_url": checkout["url"], "variant_id": checkout["variant_id"], "updated_at": _now()}, merge=True)
    _audit("lemonsqueezy_checkout_created", _identity(), tenant_id, {"payment_intent_id": intent_ref.id, "plan_id": plan_id})
    return jsonify({"created": True, "payment_intent_id": intent_ref.id, "checkout_url": checkout["url"], "plan_id": plan_id}), 201


@platform_bp.post("/client/payments/lemonsqueezy/addon-checkout")
@_require_roles("client", "operator")
def create_lemonsqueezy_addon_checkout():
    from services.lemonsqueezy_service import create_addon_checkout
    from services.plan_service import ADDONS
    payload = request.get_json(silent=True) or {}
    addon_id = str(payload.get("addon_id") or "").strip().lower()
    addon = ADDONS.get(addon_id)
    if not addon:
        return jsonify({"error": "Extra não encontrado."}), 400
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "Tenant não encontrado na sessão."}), 400
    tenant_doc = _db().collection("tenants").document(tenant_id).get()
    tenant = tenant_doc.to_dict() or {}
    if tenant.get("billing_region", "mozambique") != "international":
        return jsonify({"error": "Esta conta está configurada para M-Pesa. Usa a validação AutoPay para extras."}), 403
    intent_ref = _db().collection("payment_intents").document()
    intent_ref.set({
        "tenant_id": tenant_id,
        "provider": "lemonsqueezy",
        "payment_provider": "lemonsqueezy",
        "purchase_type": "addon",
        "addon_id": addon_id,
        "addon_name": addon["name"],
        "status": "pending_checkout",
        "created_at": _now(),
    })
    identity = _identity() or {}
    try:
        checkout = create_addon_checkout(
            addon_id=addon_id,
            tenant_id=tenant_id,
            payment_intent_id=intent_ref.id,
            email=str(identity.get("email") or tenant.get("email") or "").strip() or None,
            name=str(identity.get("name") or tenant.get("name") or "").strip() or None,
        )
    except (RuntimeError, ValueError, requests.RequestException) as exc:
        intent_ref.set({"status": "checkout_failed", "error": str(exc), "updated_at": _now()}, merge=True)
        return jsonify({"error": "O checkout do extra ainda não está disponível. Usa M-Pesa ou tenta novamente mais tarde."}), 503
    intent_ref.set({"status": "pending", "checkout_url": checkout["url"], "variant_id": checkout["variant_id"], "updated_at": _now()}, merge=True)
    _audit("lemonsqueezy_addon_checkout_created", _identity(), tenant_id, {"payment_intent_id": intent_ref.id, "addon_id": addon_id})
    return jsonify({"created": True, "payment_intent_id": intent_ref.id, "checkout_url": checkout["url"], "addon_id": addon_id}), 201


@platform_bp.post("/webhooks/lemonsqueezy")
def lemonsqueezy_webhook():
    from services.lemonsqueezy_service import event_key, extract_event, verify_signature
    raw_body = request.get_data(cache=True)
    if not verify_signature(raw_body, request.headers.get("X-Signature")):
        return jsonify({"error": "Assinatura Lemon Squeezy inválida."}), 401
    payload = request.get_json(silent=True) or {}
    event = extract_event(payload)
    key = event_key(payload)
    event_ref = _db().collection("lemonsqueezy_webhook_events").document(key)
    if event_ref.get().exists:
        return jsonify({"received": True, "duplicate": True})
    event_ref.set({"event_key": key, "event_name": event["event_name"], "object_id": event["object_id"], "received_at": _now(), "payload": payload})
    custom = event["custom_data"]
    tenant_id = str(custom.get("tenant_id") or "").strip()
    intent_id = str(custom.get("payment_intent_id") or "").strip()
    if not tenant_id or not intent_id:
        event_ref.set({"status": "unlinked"}, merge=True)
        return jsonify({"received": True, "linked": False})
    intent_ref = _db().collection("payment_intents").document(intent_id)
    intent_doc = intent_ref.get()
    intent = intent_doc.to_dict() or {}
    if not intent_doc.exists or str(intent.get("tenant_id") or "") != tenant_id:
        event_ref.set({"status": "tenant_mismatch"}, merge=True)
        return jsonify({"received": True, "linked": False})
    purchase_type = str(custom.get("purchase_type") or intent.get("purchase_type") or "plan").strip().lower()
    if purchase_type == "addon":
        from services.lemonsqueezy_service import variant_for_addon
        from services.plan_service import ADDONS
        addon_id = str(custom.get("addon_id") or intent.get("addon_id") or "").strip().lower()
        addon = ADDONS.get(addon_id)
        expected_variant = ""
        try:
            expected_variant = variant_for_addon(addon_id) if addon else ""
        except ValueError:
            expected_variant = ""
        if not addon or not expected_variant or str(event["variant_id"]) != expected_variant:
            intent_ref.set({"status": "manual_review", "provider_event": event["event_name"], "updated_at": _now()}, merge=True)
            event_ref.set({"status": "manual_review", "tenant_id": tenant_id, "payment_intent_id": intent_id}, merge=True)
            return jsonify({"received": True, "linked": True, "status": "manual_review"})
        event_name = event["event_name"]
        attributes = event["attributes"]
        active_event = event_name in {"subscription_created", "subscription_payment_success", "subscription_payment_recovered"} or (event_name == "subscription_updated" and str(attributes.get("status") or "").lower() == "active")
        if active_event:
            _db().collection("tenants").document(tenant_id).collection("addons").document(addon_id).set({
                "addon_id": addon_id,
                "name": addon["name"],
                "status": "active",
                "provider": "lemonsqueezy",
                "variant_id": expected_variant,
                "subscription_id": attributes.get("subscription_id") or event["object_id"],
                "updated_at": _now(),
            }, merge=True)
            intent_ref.set({"status": "confirmed", "provider_event": event_name, "confirmed_at": _now(), "updated_at": _now()}, merge=True)
            result_status = "confirmed"
        elif event_name in {"subscription_payment_failed", "subscription_cancelled", "subscription_expired", "order_refunded"}:
            state = "payment_failed" if event_name == "subscription_payment_failed" else ("cancelled" if "cancel" in event_name else "expired")
            _db().collection("tenants").document(tenant_id).collection("addons").document(addon_id).set({"status": state, "updated_at": _now()}, merge=True)
            intent_ref.set({"status": state, "provider_event": event_name, "updated_at": _now()}, merge=True)
            result_status = state
        else:
            intent_ref.set({"status": "received", "provider_event": event_name, "updated_at": _now()}, merge=True)
            result_status = "received"
        event_ref.set({"status": result_status, "tenant_id": tenant_id, "payment_intent_id": intent_id, "processed_at": _now()}, merge=True)
        _audit(f"lemonsqueezy_addon_{result_status}", None, tenant_id, {"event_name": event_name, "payment_intent_id": intent_id, "addon_id": addon_id, "object_id": event["object_id"]})
        return jsonify({"received": True, "linked": True, "status": result_status})
    plan_id = _lemon_plan_from_variant(event["variant_id"], str(custom.get("plan_id") or intent.get("plan_id") or ""))
    if not plan_id:
        intent_ref.set({"status": "manual_review", "provider_event": event["event_name"], "updated_at": _now()}, merge=True)
        event_ref.set({"status": "manual_review", "tenant_id": tenant_id, "payment_intent_id": intent_id}, merge=True)
        return jsonify({"received": True, "linked": True, "status": "manual_review"})
    event_name = event["event_name"]
    attributes = event["attributes"]
    if event_name in {"subscription_created", "subscription_payment_success", "subscription_payment_recovered"} or (event_name == "subscription_updated" and str(attributes.get("status") or "").lower() == "active"):
        _apply_lemon_plan(tenant_id, plan_id, attributes, intent_ref, event)
        result_status = "confirmed"
    elif event_name in {"subscription_payment_failed"}:
        intent_ref.set({"status": "payment_failed", "provider_event": event_name, "updated_at": _now()}, merge=True)
        result_status = "payment_failed"
    elif event_name in {"subscription_cancelled", "subscription_expired", "order_refunded"}:
        intent_ref.set({"status": "cancelled" if "cancel" in event_name else "expired", "provider_event": event_name, "updated_at": _now()}, merge=True)
        _db().collection("tenants").document(tenant_id).set({"status_plano": "cancelado" if "cancel" in event_name else "expirado", "status": "cancelado" if "cancel" in event_name else "expirado", "updated_at": _now()}, merge=True)
        result_status = "cancelled" if "cancel" in event_name else "expired"
    else:
        intent_ref.set({"status": "received", "provider_event": event_name, "updated_at": _now()}, merge=True)
        result_status = "received"
    event_ref.set({"status": result_status, "tenant_id": tenant_id, "payment_intent_id": intent_id, "processed_at": _now()}, merge=True)
    _audit(f"lemonsqueezy_{result_status}", None, tenant_id, {"event_name": event_name, "payment_intent_id": intent_id, "object_id": event["object_id"]})
    return jsonify({"received": True, "linked": True, "status": result_status})


@platform_bp.get("/admin/audit")
@_require_roles("owner", "admin")
def list_audit_events():
    rows = []
    for document in _db().collection("audit_events").limit(200).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    return jsonify({"events": rows})


@platform_bp.get("/admin/health")
@_require_roles("owner", "admin")
def admin_health():
    services = {"backend": "online", "firestore": "online", "redis": "unknown", "evolution": "unknown"}
    try:
        import redis
        redis.from_url(os.getenv("REDIS_URL", "redis://redis:6379/1"), decode_responses=True).ping()
        services["redis"] = "online"
    except Exception:
        services["redis"] = "offline"
    base_url = str(os.getenv("EVOLUTION_API_URL", "")).rstrip("/")
    if base_url and os.getenv("EVOLUTION_API_KEY"):
        try:
            response = requests.get(f"{base_url}/instance/fetchInstances", headers={"apikey": os.getenv("EVOLUTION_API_KEY")}, timeout=8)
            services["evolution"] = "online" if response.ok else "degraded"
        except requests.RequestException:
            services["evolution"] = "offline"
    return jsonify({"services": services, "worker": "managed by Docker Compose"})


@platform_bp.get("/client/plans")
@_require_roles("client", "operator")
def client_plans_catalog():
    from services.plan_service import ADDONS
    return jsonify({
        "plans": public_plan_rows(),
        "addons": [{"id": addon_id, **data} for addon_id, data in ADDONS.items()],
        "trial_days": 2,
        "mpesa_number": "855000929",
        "mpesa_name": "Abel Francisco",
    })


@platform_bp.post("/client/evolution/qr")
@_require_roles("client", "operator")
def request_evolution_qr():
    tenant_id = _tenant_for_identity(_identity())
    payload = request.get_json(silent=True) or {}
    tenant_ref = _db().collection("tenants").document(tenant_id)
    tenant = tenant_ref.get().to_dict() or {}
    status = str(tenant.get("status_plano", tenant.get("status", "demonstracao"))).lower()
    if status in {"expirado", "suspenso", "cancelado"} or _plan_expired(tenant):
        return jsonify({"error": "A demonstração/plano expirou. Ativa ou renova o plano antes de ligar o WhatsApp ou gerar outro QR Code."}), 402
    phone = re.sub(r"\D", "", str(payload.get("phone") or tenant.get("telefone_proprietario") or tenant.get("phone") or ""))
    if len(phone) < 8:
        return jsonify({"error": "Indica o número de WhatsApp que será automatizado."}), 400
    try:
        from services.evolution_service import criar_e_configurar_instancia_automatica, obter_qrcode_instancia
        if not criar_e_configurar_instancia_automatica(phone):
            return jsonify({"error": "Não foi possível preparar a instância Evolution."}), 502
        data = obter_qrcode_instancia(phone)
        instance_name = data["instance_name"]
        state = data["state"]
        base64_qr = data.get("base64")
        trial_fields = {}
        if not is_paid_plan(tenant):
            trial_fields = active_fields(phone) if state == "open" else pending_fields(phone)
        tenant_ref.set({
            **trial_fields,
            "instance_name": instance_name,
            "telefone_proprietario": phone,
            "evolution_state": state,
            "updated_at": _now(),
        }, merge=True)
        profile_sync = {key: tenant.get(key) for key in ("empresa_nome", "nicho", "email_corporativo", "redes_sociais") if tenant.get(key) is not None}
        if profile_sync:
            _db().collection("clientes_bot").document(instance_name).set(profile_sync, merge=True)
        return jsonify({"state": state, "instance_name": instance_name, "qrcode": base64_qr and (base64_qr if str(base64_qr).startswith("data:") else f"data:image/png;base64,{base64_qr}")})
    except requests.RequestException:
        return jsonify({"error": "A Evolution API não respondeu ao pedido de QR Code."}), 502
    except Exception:
        return jsonify({"error": "Não foi possível gerar o QR Code neste momento."}), 502


_PUBLIC_CHAT_RATE: dict[str, list[float]] = {}


def _public_plan_answer() -> str:
    return (
        "Aqui estão os planos NEGOBOT-MOZ:\n\n"
        "Para clientes em Moçambique:\n"
        "• Básico — 500 MT/mês: até 1.500 conversas/contactos, 1 utilizador, 2 campanhas por mês e WhatsApp.\n"
        "• Médio — 1.000 MT/mês: até 5.000 conversas/contactos, 3 utilizadores, 10 campanhas por mês e mais 1 canal aprovado.\n"
        "• Premium — 1.500 MT/mês: até 15.000 conversas/contactos, 5 utilizadores, 25 campanhas por mês e até 3 canais adicionais aprovados.\n\n"
        "Para clientes internacionais, os mesmos planos são apresentados em USD: Basic USD 8/mês, Growth USD 16/mês e Premium USD 24/mês. O pagamento internacional é feito por cartão ou PayPal através do checkout Lemon Squeezy. Os preços em USD são valores comerciais fixos; o checkout mostra o valor final configurado.\n\n"
        "Todos os planos têm validade de 30 dias e começam com demonstração de 2 dias. Clientes de Moçambique pagam por M-Pesa para 855000929, em nome de Abel Francisco; depois da validação pelo AutoPay, a Evolution API prepara o QR Code para ligar o WhatsApp."
    )


def _public_is_plan_question(message: str) -> bool:
    normalized = message.casefold()
    return any(term in normalized for term in ("plano", "preço", "preco", "quanto", "custa", "mensal", "mt", "benefício", "beneficio"))


@platform_bp.post("/public/assistant/chat")
def public_assistant_chat():
    payload = request.get_json(silent=True) or {}
    message = str(payload.get("message") or "").strip()
    source = str(payload.get("source") or "platform").lower()
    if source not in {"platform", "facebook", "instagram", "whatsapp", "other"}:
        source = "platform"
    if not message or len(message) > 1200:
        return jsonify({"error": "Escreve uma mensagem entre 1 e 1.200 caracteres."}), 400
    now = time.time()
    visitor = request.headers.get("X-Forwarded-For", request.remote_addr or "public").split(",")[0].strip()
    recent = [stamp for stamp in _PUBLIC_CHAT_RATE.get(visitor, []) if now - stamp < 600]
    if len(recent) >= 12:
        return jsonify({"error": "Atingiste o limite temporário de mensagens. Tenta novamente mais tarde."}), 429
    recent.append(now)
    _PUBLIC_CHAT_RATE[visitor] = recent
    try:
        _db().collection("public_leads").add({"source": source, "message": message[:1200], "created_at": _now()})
    except Exception:
        pass
    prompt = """És o assistente comercial público do NEGOBOT-MOZ, em Português de Moçambique. Explica apenas os planos reais: para Moçambique, Básico 500 MT com até 1.500 conversas/contactos, Médio 1.000 MT com até 5.000 e Premium 1.500 MT com até 15.000; para clientes internacionais, mostra USD 8, USD 16 e USD 24, respectivamente. Todos têm validade de 30 dias e demonstração de 2 dias. M-Pesa é apenas para Moçambique; clientes internacionais usam cartão ou PayPal no checkout Lemon Squeezy. Nunca digas que um pagamento foi confirmado sem validação. Sê comercial, honesto e breve. Não inventes preços, limites ou integrações."""
    if _public_is_plan_question(message):
        answer = _public_plan_answer()
    else:
        try:
            public_tenant = "public:" + hashlib.sha256(visitor.encode("utf-8")).hexdigest()[:24]
            ai_result = request_ai_text(
                tenant_id=public_tenant,
                messages=[{"role": "user", "content": message}],
                system_prompt=prompt,
            )
            answer = str(ai_result.get("text") or "").strip()
        except AIQueueError:
            answer = "Posso ajudar com os planos, pagamentos M-Pesa, ligação do WhatsApp e demonstração de 2 dias."
        if not answer or "processar muitas mensagens" in answer.casefold():
            answer = "Posso ajudar com os planos, pagamentos M-Pesa, ligação do WhatsApp e demonstração de 2 dias. Escreve 'planos' para veres a tabela completa."
    return jsonify({"answer": answer, "source": source, "next": {"whatsapp": "/falar-whatsapp", "platform": "/plataforma"}})


# ─────────────────────────────────────────────────────────────────────────────
# Métricas, relatórios e suporte SaaS
# ─────────────────────────────────────────────────────────────────────────────

_TICKET_STATUSES = {"open", "in_progress", "waiting_client", "resolved", "closed"}
_TICKET_PRIORITIES = {"low", "normal", "high", "urgent"}


def _tenant_collection_rows(collection: str, tenant_id: str, limit: int = 2000) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for document in _db().collection(collection).where("tenant_id", "==", tenant_id).limit(limit).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    return rows


def _campaign_metrics_for_tenant(tenant_id: str) -> dict[str, Any]:
    campaigns = _tenant_collection_rows("campaigns", tenant_id, 500)
    recipients = _tenant_collection_rows("campaign_recipients", tenant_id, 5000)
    conversations = list(_db().collection("clientes_bot").document(tenant_id).collection("conversas").limit(500).stream())
    contacts = _tenant_collection_rows("contacts", tenant_id, 5000)
    statuses: dict[str, int] = {}
    for campaign in campaigns:
        status = str(campaign.get("status") or "unknown")
        statuses[status] = statuses.get(status, 0) + 1
    recipient_statuses: dict[str, int] = {}
    for recipient in recipients:
        status = str(recipient.get("status") or "unknown")
        recipient_statuses[status] = recipient_statuses.get(status, 0) + 1
    sent = recipient_statuses.get("sent", 0) + recipient_statuses.get("delivered", 0)
    failed = recipient_statuses.get("failed", 0)
    attempted = sent + failed
    recent = sorted(campaigns, key=lambda item: str(item.get("created_at") or ""), reverse=True)[:10]
    return {
        "contacts": {"total": len(contacts), "opt_in": sum(1 for item in contacts if item.get("opt_in", True)), "opt_out": sum(1 for item in contacts if not item.get("opt_in", True))},
        "conversations": len(conversations),
        "campaigns": {"total": len(campaigns), "by_status": statuses, "recent": recent},
        "deliveries": {"total": len(recipients), "by_status": recipient_statuses, "sent": sent, "failed": failed, "delivery_rate": round((sent / attempted) * 100, 2) if attempted else 0},
    }


@platform_bp.get("/client/metrics")
@_require_roles("client", "operator")
def client_metrics():
    tenant_id = _tenant_for_identity(_identity())
    return jsonify({"tenant_id": tenant_id, "metrics": _campaign_metrics_for_tenant(tenant_id)})


@platform_bp.get("/client/reports/campaigns")
@_require_roles("client", "operator")
def client_campaign_report():
    tenant_id = _tenant_for_identity(_identity())
    metrics = _campaign_metrics_for_tenant(tenant_id)
    return jsonify({"tenant_id": tenant_id, "generated_at": _now(), "campaigns": metrics["campaigns"], "deliveries": metrics["deliveries"]})


@platform_bp.get("/client/support/tickets")
@_require_roles("client", "operator")
def client_support_tickets():
    tenant_id = _tenant_for_identity(_identity())
    rows = _tenant_collection_rows("support_tickets", tenant_id, 200)
    rows.sort(key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True)
    return jsonify({"tickets": rows})


@platform_bp.post("/client/support/tickets")
@_require_roles("client", "operator")
def create_support_ticket():
    payload = request.get_json(silent=True) or {}
    subject = str(payload.get("subject") or "").strip()
    message = str(payload.get("message") or "").strip()
    category = str(payload.get("category") or "general").strip().lower()[:40]
    priority = str(payload.get("priority") or "normal").strip().lower()
    if not 4 <= len(subject) <= 160 or not 10 <= len(message) <= 5000:
        return jsonify({"error": "Indica um assunto e uma descrição válidos."}), 400
    if priority not in _TICKET_PRIORITIES:
        priority = "normal"
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("support_tickets").document()
    now = _now()
    ticket = {"tenant_id": tenant_id, "subject": subject, "message": message, "category": category or "general", "priority": priority, "status": "open", "created_by": (_identity() or {}).get("id"), "created_at": now, "updated_at": now}
    reference.set(ticket)
    _audit("support_ticket_created", _identity(), tenant_id, {"ticket_id": reference.id, "priority": priority, "category": category})
    return jsonify({"created": True, "ticket": {"id": reference.id, **ticket}}), 201


@platform_bp.patch("/client/support/tickets/<ticket_id>")
@_require_roles("client", "operator")
def update_client_support_ticket(ticket_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("support_tickets").document(ticket_id)
    document = reference.get()
    data = document.to_dict() if document.exists else None
    if not data or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Ticket não encontrado."}), 404
    payload = request.get_json(silent=True) or {}
    changes: dict[str, Any] = {}
    if "message" in payload:
        message = str(payload.get("message") or "").strip()
        if not 10 <= len(message) <= 5000:
            return jsonify({"error": "A mensagem deve ter entre 10 e 5000 caracteres."}), 400
        changes["last_client_message"] = message
        changes["last_client_message_at"] = _now()
        if data.get("status") == "waiting_client":
            changes["status"] = "open"
    if "status" in payload and str(payload.get("status")) in {"closed", "open"}:
        changes["status"] = str(payload["status"])
    if not changes:
        return jsonify({"error": "Nenhuma alteração válida foi enviada."}), 400
    changes["updated_at"] = _now()
    reference.set(changes, merge=True)
    _audit("support_ticket_updated", _identity(), tenant_id, {"ticket_id": ticket_id, "fields": sorted(changes)})
    return jsonify({"updated": True, "ticket_id": ticket_id, "fields": sorted(changes)})


@platform_bp.get("/admin/metrics")
@_require_roles("owner", "admin")
def admin_metrics():
    db = _db()
    tenants = list(db.collection("tenants").limit(5000).stream())
    users = list(db.collection("platform_users").limit(5000).stream())
    campaigns = list(db.collection("campaigns").limit(5000).stream())
    payments = list(db.collection("payment_intents").limit(5000).stream())
    tickets = list(db.collection("support_tickets").limit(5000).stream())
    return jsonify({
        "generated_at": _now(),
        "tenants": {"total": len(tenants), "active": sum(1 for item in tenants if (item.to_dict() or {}).get("status", "active") in {"active", "ativo"})},
        "users": {"total": len(users), "active": sum(1 for item in users if (item.to_dict() or {}).get("status", "active") == "active")},
        "campaigns": {"total": len(campaigns), "by_status": _count_statuses(campaigns)},
        "payments": {"total": len(payments), "confirmed": sum(1 for item in payments if (item.to_dict() or {}).get("status") in {"confirmed", "active"})},
        "support": {"total": len(tickets), "open": sum(1 for item in tickets if (item.to_dict() or {}).get("status", "open") in {"open", "in_progress", "waiting_client"})},
    })


def _count_statuses(documents: list[Any]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for document in documents:
        status = str((document.to_dict() or {}).get("status") or "unknown")
        counts[status] = counts.get(status, 0) + 1
    return counts


@platform_bp.get("/admin/support/tickets")
@_require_roles("owner", "admin")
def admin_support_tickets():
    rows = []
    for document in _db().collection("support_tickets").limit(500).stream():
        item = document.to_dict() or {}
        item["id"] = document.id
        rows.append(item)
    rows.sort(key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True)
    return jsonify({"tickets": rows})


@platform_bp.patch("/admin/support/tickets/<ticket_id>")
@_require_roles("owner", "admin")
def admin_update_support_ticket(ticket_id: str):
    reference = _db().collection("support_tickets").document(ticket_id)
    document = reference.get()
    if not document.exists:
        return jsonify({"error": "Ticket não encontrado."}), 404
    payload = request.get_json(silent=True) or {}
    status = str(payload.get("status") or "").strip().lower()
    if status not in _TICKET_STATUSES:
        return jsonify({"error": "Estado de ticket inválido."}), 400
    changes = {"status": status, "updated_at": _now(), "updated_by": (_identity() or {}).get("id")}
    if "reply" in payload:
        reply = str(payload.get("reply") or "").strip()
        if reply:
            changes["last_admin_reply"] = reply[:5000]
            changes["last_admin_reply_at"] = _now()
    reference.set(changes, merge=True)
    tenant_id = str((document.to_dict() or {}).get("tenant_id") or "") or None
    _audit("support_ticket_admin_updated", _identity(), tenant_id, {"ticket_id": ticket_id, "status": status})
    return jsonify({"updated": True, "ticket_id": ticket_id, "status": status})


_VIDEO_ASSET_EXTENSIONS = {
    ".mp4": ("video", "video/mp4"),
    ".mov": ("video", "video/quicktime"),
    ".webm": ("video", "video/webm"),
    ".png": ("image", "image/png"),
    ".jpg": ("image", "image/jpeg"),
    ".jpeg": ("image", "image/jpeg"),
    ".mp3": ("audio", "audio/mpeg"),
    ".wav": ("audio", "audio/wav"),
    ".ogg": ("audio", "audio/ogg"),
    ".m4a": ("audio", "audio/mp4"),
}
_VIDEO_ASSET_MAX_BYTES = 16 * 1024 * 1024


def _video_asset_url(asset_id: str) -> str:
    base_url = str(os.getenv("VIDEO_ASSET_BASE_URL") or os.getenv("PUBLIC_APP_BASE_URL") or request.host_url).rstrip("/")
    return f"{base_url}/api/platform/client/videos/assets/{quote(asset_id)}"


def _serialise_video_asset(asset_id: str, data: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": asset_id,
        "file_name": str(data.get("file_name") or "media"),
        "size_bytes": int(data.get("size_bytes") or 0),
        "mime_type": str(data.get("mime_type") or "application/octet-stream"),
        "kind": str(data.get("kind") or "video"),
        "asset_url": _video_asset_url(asset_id),
        "created_at": data.get("created_at"),
    }


@platform_bp.get("/client/videos/assets")
@_require_roles("client", "operator")
def list_video_assets():
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "tenant não configurado"}), 403
    documents = _db().collection("video_assets").where("tenant_id", "==", tenant_id).limit(100).stream()
    assets = [_serialise_video_asset(document.id, document.to_dict() or {}) for document in documents]
    assets.sort(key=lambda item: str(item.get("created_at") or ""), reverse=True)
    return jsonify({"assets": assets, "count": len(assets)})


@platform_bp.post("/client/videos/assets")
@_require_tenant_roles("owner", "operator")
def upload_video_asset():
    tenant_id = _tenant_for_identity(_identity())
    if not tenant_id:
        return jsonify({"error": "tenant não configurado"}), 403
    uploaded = request.files.get("file")
    filename = secure_filename(uploaded.filename if uploaded else "")
    extension = (os.path.splitext(filename)[1] or "").lower()
    if uploaded is None or not filename:
        return jsonify({"error": "Selecciona um ficheiro de media."}), 400
    if extension not in _VIDEO_ASSET_EXTENSIONS:
        return jsonify({"error": "Formato não suportado. Usa MP4, MOV, WEBM, PNG, JPG, JPEG, MP3, WAV, OGG ou M4A."}), 400
    content = uploaded.read(_VIDEO_ASSET_MAX_BYTES + 1)
    if len(content) > _VIDEO_ASSET_MAX_BYTES:
        return jsonify({"error": "O ficheiro excede o limite de 16 MB."}), 413
    if not content:
        return jsonify({"error": "O ficheiro está vazio."}), 400
    kind, fallback_mime = _VIDEO_ASSET_EXTENSIONS[extension]
    mime_type = str(uploaded.mimetype or fallback_mime).lower()
    if extension == ".webm" and mime_type.startswith("audio/"):
        kind, fallback_mime = "audio", "audio/webm"
    if not (mime_type.startswith(f"{kind}/") or mime_type == fallback_mime):
        mime_type = fallback_mime
    asset_id = secrets.token_urlsafe(18)
    storage_key = store_blob(tenant_id, asset_id, filename, content, mime_type, prefix="video-assets")
    if not storage_key:
        return jsonify({"error": "Não foi possível guardar o ficheiro de media."}), 503
    data = {"tenant_id": tenant_id, "file_name": filename, "size_bytes": len(content), "mime_type": mime_type, "kind": kind, "storage_key": storage_key, "created_at": _now()}
    _db().collection("video_assets").document(asset_id).set(data)
    _audit("video_asset_uploaded", _identity(), tenant_id, {"asset_id": asset_id, "kind": kind, "size_bytes": len(content)})
    return jsonify({"uploaded": True, "asset": _serialise_video_asset(asset_id, data)}), 201


@platform_bp.delete("/client/videos/assets/<asset_id>")
@_require_tenant_roles("owner", "operator")
def delete_video_asset(asset_id: str):
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("video_assets").document(asset_id)
    document = reference.get()
    data = document.to_dict() if document.exists else {}
    if not document.exists or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Media não encontrado neste tenant."}), 404
    delete_original(data.get("storage_key"))
    reference.delete()
    _audit("video_asset_deleted", _identity(), tenant_id, {"asset_id": asset_id})
    return jsonify({"deleted": True, "asset_id": asset_id})


@platform_bp.get("/client/videos/assets/<asset_id>")
def stream_video_asset(asset_id: str):
    expected_token = str(os.getenv("VIDEO_SERVICE_TOKEN") or "").strip()
    supplied_token = str(request.headers.get("X-Video-Service-Token") or "").strip()
    if not expected_token or not hmac.compare_digest(supplied_token, expected_token):
        return jsonify({"error": "Media não encontrado."}), 404
    tenant_id = str(request.headers.get("X-Video-Tenant-Id") or "").strip()
    document = _db().collection("video_assets").document(asset_id).get()
    data = document.to_dict() if document.exists else {}
    if not document.exists or not tenant_id or data.get("tenant_id") != tenant_id:
        return jsonify({"error": "Media não encontrado."}), 404
    content = read_blob(data.get("storage_key"))
    if content is None:
        return jsonify({"error": "Media não disponível."}), 404
    return Response(content, mimetype=str(data.get("mime_type") or "application/octet-stream"), headers={"Cache-Control": "private, no-store", "Content-Disposition": f'inline; filename="{secure_filename(str(data.get("file_name") or "media"))}"'})


@platform_bp.post("/client/videos/jobs")
@_require_roles("client", "operator")
def create_video_job():
    base_url = str(os.getenv("VIDEO_SERVICE_URL", "")).rstrip("/")
    service_token = os.getenv("VIDEO_SERVICE_TOKEN", "").strip()
    if not base_url or not service_token:
        return jsonify({"error": "O motor de vídeos ainda não está configurado."}), 503
    payload = request.get_json(silent=True) or {}
    scenes = payload.get("scenes") or []
    title = str(payload.get("title") or "").strip()
    if not 2 <= len(title) <= 160 or not isinstance(scenes, list) or not 1 <= len(scenes) <= 20:
        return jsonify({"error": "Indica um título e pelo menos uma cena válida."}), 400
    total_characters = len(str(payload.get("narracao") or ""))
    total_duration = 0.0
    for scene in scenes:
        if not isinstance(scene, dict):
            return jsonify({"error": "Cada cena deve ser um objecto válido."}), 422
        text = str(scene.get("text") or "")
        total_characters += len(text)
        try:
            requested_duration = float(scene.get("duration_seconds") or 3.5)
        except (TypeError, ValueError):
            return jsonify({"error": "A duração de cada cena deve ser numérica."}), 422
        words = len(text.split())
        narration_duration = words / 2.35 if words else 0.0
        total_duration += max(requested_duration, narration_duration)
    if total_characters > _VIDEO_MAX_TOTAL_SCRIPT_CHARACTERS:
        return jsonify({"error": "O roteiro completo não pode ultrapassar 5.000 caracteres."}), 422
    if total_duration > _VIDEO_MAX_TOTAL_DURATION_SECONDS:
        return jsonify({"error": "A duração total calculada não pode ultrapassar 300 segundos (5 minutos)."}), 422
    tenant_id = _tenant_for_identity(_identity())
    outgoing = {"tenant_id": tenant_id, "title": title, "scenes": scenes, "language": str(payload.get("language") or "pt-MZ"), "voice": payload.get("voice"), "subtitles": bool(payload.get("subtitles", True)), "narracao": payload.get("narracao"), "palavras_chave": payload.get("palavras_chave") or [], "background_keywords": payload.get("background_keywords") or [], "transition": str(payload.get("transition") or "fade")}
    try:
        response = requests.post(f"{base_url}/api/video/jobs", json=outgoing, headers={"X-Video-Service-Token": service_token}, timeout=20)
    except requests.RequestException:
        return jsonify({"error": "O motor de vídeos está temporariamente indisponível."}), 503
    if not response.ok:
        detail = "O motor de vídeos rejeitou o job."
        if response.status_code == 422:
            try:
                raw_detail = response.json().get("detail")
                if isinstance(raw_detail, list):
                    messages = [str(item.get("msg") or "Validação inválida.") for item in raw_detail if isinstance(item, dict)]
                    detail = " ".join(messages) or detail
                elif raw_detail:
                    detail = str(raw_detail)
            except (ValueError, TypeError):
                pass
            return jsonify({"error": detail}), 422
        return jsonify({"error": detail}), 502
    data = response.json()
    job = data.get("job") or {}
    job_id = str(job.get("id") or "").strip()
    if not job_id:
        return jsonify({"error": "O motor de vídeos devolveu uma resposta inválida."}), 502
    _db().collection("video_jobs").document(job_id).set({"tenant_id": tenant_id, "job_id": job_id, "title": title, "status": "queued", "created_at": _now()}, merge=True)
    _audit("video_job_created", _identity(), tenant_id, {"job_id": job_id, "scenes": len(scenes)})
    return jsonify({"accepted": True, "job": job}), 202


@platform_bp.get("/client/videos/jobs/<job_id>")
@_require_roles("client", "operator")
def get_video_job(job_id: str):
    base_url = str(os.getenv("VIDEO_SERVICE_URL", "")).rstrip("/")
    service_token = os.getenv("VIDEO_SERVICE_TOKEN", "").strip()
    tenant_id = _tenant_for_identity(_identity())
    document = _db().collection("video_jobs").document(job_id).get()
    if not document.exists or (document.to_dict() or {}).get("tenant_id") != tenant_id:
        return jsonify({"error": "Job de vídeo não encontrado."}), 404
    if not base_url or not service_token:
        return jsonify({"error": "O motor de vídeos ainda não está configurado."}), 503
    try:
        response = requests.get(f"{base_url}/api/video/jobs/{quote(job_id)}", headers={"X-Video-Service-Token": service_token, "X-Video-Tenant-Id": tenant_id}, timeout=15)
    except requests.RequestException:
        return jsonify({"error": "O motor de vídeos está temporariamente indisponível."}), 503
    if not response.ok:
        return jsonify({"error": "Não foi possível obter o estado do vídeo."}), 502
    data = response.json()
    job = data.get("job") or {}
    _db().collection("video_jobs").document(job_id).set({"status": job.get("status"), "progress": job.get("progress", 0), "updated_at": _now(), "output_available": bool(job.get("output_available"))}, merge=True)
    return jsonify({"job": job})


@platform_bp.get("/client/videos/jobs/<job_id>/preview")
@_require_roles("client", "operator")
def preview_video_job(job_id: str):
    base_url = str(os.getenv("VIDEO_SERVICE_URL", "")).rstrip("/")
    service_token = os.getenv("VIDEO_SERVICE_TOKEN", "").strip()
    tenant_id = _tenant_for_identity(_identity())
    document = _db().collection("video_jobs").document(job_id).get()
    if not document.exists or (document.to_dict() or {}).get("tenant_id") != tenant_id:
        return jsonify({"error": "Job de vídeo não encontrado."}), 404
    if not base_url or not service_token:
        return jsonify({"error": "O motor de vídeos ainda não está configurado."}), 503
    upstream_headers = {
        "X-Video-Service-Token": service_token,
        "X-Video-Tenant-Id": tenant_id,
    }
    if request.headers.get("Range"):
        upstream_headers["Range"] = request.headers["Range"]
    try:
        upstream = requests.get(
            f"{base_url}/api/video/jobs/{quote(job_id)}/preview",
            headers=upstream_headers,
            stream=True,
            timeout=(10, 120),
        )
    except requests.RequestException:
        return jsonify({"error": "O motor de vídeos está temporariamente indisponível."}), 503
    if not upstream.ok:
        try:
            detail = (upstream.json() or {}).get("detail")
        except ValueError:
            detail = None
        upstream.close()
        return jsonify({"error": detail or "O vídeo ainda não está disponível para pré-visualização."}), upstream.status_code if upstream.status_code in {404, 409} else 502

    def relay_preview():
        try:
            for chunk in upstream.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    yield chunk
        finally:
            upstream.close()

    headers = {"Cache-Control": "no-store", "X-Accel-Buffering": "no"}
    for key in ("Content-Length", "Content-Range", "Accept-Ranges", "Content-Disposition"):
        if upstream.headers.get(key):
            headers[key] = upstream.headers[key]
    return Response(
        stream_with_context(relay_preview()),
        status=upstream.status_code,
        mimetype=upstream.headers.get("Content-Type", "video/mp4"),
        headers=headers,
    )


@platform_bp.get("/client/videos/jobs/<job_id>/download")
@_require_roles("client", "operator")
def download_video_job(job_id: str):
    base_url = str(os.getenv("VIDEO_SERVICE_URL", "")).rstrip("/")
    service_token = os.getenv("VIDEO_SERVICE_TOKEN", "").strip()
    tenant_id = _tenant_for_identity(_identity())
    document = _db().collection("video_jobs").document(job_id).get()
    if not document.exists or (document.to_dict() or {}).get("tenant_id") != tenant_id:
        return jsonify({"error": "Job de vídeo não encontrado."}), 404
    if not base_url or not service_token:
        return jsonify({"error": "O motor de vídeos ainda não está configurado."}), 503
    try:
        upstream = requests.get(
            f"{base_url}/api/video/jobs/{quote(job_id)}/download",
            headers={"X-Video-Service-Token": service_token, "X-Video-Tenant-Id": tenant_id},
            stream=True,
            timeout=(10, 120),
        )
    except requests.RequestException:
        return jsonify({"error": "O motor de vídeos está temporariamente indisponível."}), 503
    if not upstream.ok:
        try:
            detail = (upstream.json() or {}).get("detail")
        except ValueError:
            detail = None
        upstream.close()
        return jsonify({"error": detail or "O vídeo já não está disponível para download."}), upstream.status_code if upstream.status_code in {404, 409} else 502

    def relay():
        try:
            for chunk in upstream.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    yield chunk
        finally:
            upstream.close()

    headers = {"Cache-Control": "no-store", "X-Accel-Buffering": "no"}
    for key in ("Content-Length", "Content-Disposition"):
        if upstream.headers.get(key):
            headers[key] = upstream.headers[key]
    return Response(stream_with_context(relay()), status=200, mimetype=upstream.headers.get("Content-Type", "video/mp4"), headers=headers)


@platform_bp.delete("/client/videos/jobs/<job_id>")
@_require_roles("client", "operator")
def delete_video_job(job_id: str):
    base_url = str(os.getenv("VIDEO_SERVICE_URL", "")).rstrip("/")
    service_token = os.getenv("VIDEO_SERVICE_TOKEN", "").strip()
    tenant_id = _tenant_for_identity(_identity())
    document = _db().collection("video_jobs").document(job_id).get()
    if not document.exists or (document.to_dict() or {}).get("tenant_id") != tenant_id:
        return jsonify({"error": "Job de vídeo não encontrado."}), 404
    if not base_url or not service_token:
        return jsonify({"error": "O motor de vídeos ainda não está configurado."}), 503
    try:
        response = requests.delete(f"{base_url}/api/video/jobs/{quote(job_id)}", headers={"X-Video-Service-Token": service_token, "X-Video-Tenant-Id": tenant_id}, timeout=20)
    except requests.RequestException:
        return jsonify({"error": "O motor de vídeos está temporariamente indisponível."}), 503
    if not response.ok:
        try:
            detail = (response.json() or {}).get("detail")
        except ValueError:
            detail = None
        return jsonify({"error": detail or "Não foi possível apagar o vídeo."}), response.status_code if response.status_code in {404, 409} else 502
    _db().collection("video_jobs").document(job_id).set({"status": "deleted", "output_available": False, "deleted_at": _now(), "deletion_reason": "manual"}, merge=True)
    _audit("video_job_deleted", _identity(), tenant_id, {"job_id": job_id, "reason": "manual"})
    return jsonify({"deleted": True, "job_id": job_id})


@platform_bp.get("/client/channels")
@_require_roles("client", "operator")
def client_channels():
    tenant_id = _tenant_for_identity(_identity())
    tenant_document = _db().collection("tenants").document(tenant_id).get()
    if not tenant_document.exists:
        return jsonify({"error": "Tenant não encontrado."}), 404
    tenant = tenant_document.to_dict() or {}
    return jsonify({"tenant_id": tenant_id, "channels": client_channel_rows(tenant)})


@platform_bp.get("/client/channels/telegram")
@_require_roles("client", "operator")
def telegram_channel_status():
    tenant_id = _tenant_for_identity(_identity())
    tenant = _tenant_data(tenant_id)
    telegram = dict((tenant.get("channels") or {}).get("telegram") or {})
    info = telegram.get("last_webhook_info") if isinstance(telegram.get("last_webhook_info"), dict) else {}
    return jsonify({
        "channel": "telegram",
        "status": telegram.get("status", "not_configured"),
        "bot": {"id": telegram.get("bot_id"), "username": telegram.get("bot_username"), "name": telegram.get("bot_name")},
        "webhook_url": telegram.get("webhook_url"),
        "last_event_at": telegram.get("last_event_at"),
        "last_error": telegram.get("last_error") or info.get("last_error_message"),
        "pending_update_count": info.get("pending_update_count", 0),
        "has_token": bool(telegram.get("token_ciphertext")),
    })


@platform_bp.post("/client/channels/telegram/connect")
@_require_tenant_roles("owner", "operator")
def connect_telegram():
    payload = request.get_json(silent=True) or {}
    token = str(payload.get("bot_token") or "").strip()
    if not token or len(token) > 256:
        return jsonify({"error": "Introduz um token Telegram válido."}), 400
    tenant_id = _tenant_for_identity(_identity())
    webhook_base = str(os.getenv("PUBLIC_API_BASE_URL") or "https://negobot-api.duckdns.org").rstrip("/")
    webhook_url = f"{webhook_base}/api/omnichannel/telegram/{tenant_id}"
    try:
        bot = get_me(token)
        secret_token = secrets.token_hex(32)
        set_webhook(token, url=webhook_url, secret_token=secret_token)
        info = get_webhook_info(token)
    except (TelegramApiError, SecretStoreError) as exc:
        _audit("telegram_webhook_setup_failed", _identity(), tenant_id, {"error": str(exc)[:240]})
        return jsonify({"error": "Não foi possível validar ou registar o bot Telegram."}), 502
    if str(info.get("url") or "") != webhook_url:
        _audit("telegram_webhook_verification_failed", _identity(), tenant_id, {})
        return jsonify({"error": "O Telegram não confirmou o URL do webhook."}), 502
    tenant_ref = _db().collection("tenants").document(tenant_id)
    tenant = tenant_ref.get().to_dict() or {}
    central_account_id = central_account_id_for_tenant(tenant)
    registry = registry_status(_db(), central_account_id)
    if not is_paid_plan(tenant):
        if central_account_id:
            if registry_is_expired(registry) or str(registry.get("trial_status") or "").lower() == "trial_expired":
                return jsonify({"error": "A demonstração terminou. Escolhe um plano pago para ligar outro canal."}), 403
            if not registry.get("trial_consumed"):
                claimed, registry = claim_trial_for_account(_db(), central_account_id, tenant_id, "telegram", email=tenant.get("account_email") or tenant.get("email"))
                if not claimed and registry.get("blocked_by_identity"):
                    return jsonify({"error": "Esta identidade já utilizou a demonstração. Escolhe um plano pago para ligar este canal."}), 403
                if not claimed:
                    registry = registry_status(_db(), central_account_id)
            tenant_trial_fields = trial_fields_from_registry(registry, "telegram", bot.get("username") or str(bot.get("id")))
        else:
            tenant_trial_fields = active_fields(str(tenant.get("instance_name") or tenant_id))
    else:
        tenant_trial_fields = {}
    channels = dict(tenant.get("channels") or {}) if isinstance(tenant.get("channels"), dict) else {}
    channels["telegram"] = {
        "status": "connected",
        "setup": "bot_token",
        "provider": "Telegram Bot API",
        "bot_id": str(bot.get("id")),
        "bot_username": bot.get("username"),
        "bot_name": bot.get("first_name"),
        "token_ciphertext": encrypt_secret(token),
        "webhook_secret_ciphertext": encrypt_secret(secret_token),
        "webhook_url": webhook_url,
        "last_webhook_info": {"pending_update_count": info.get("pending_update_count", 0), "last_error_message": info.get("last_error_message"), "last_error_date": info.get("last_error_date")},
        "connected_at": _now(),
        "updated_at": _now(),
    }
    tenant_ref.set({"channels": channels, **tenant_trial_fields, "central_account_id": central_account_id, "updated_at": _now()}, merge=True)
    _audit("telegram_webhook_connected", _identity(), tenant_id, {"bot_id": str(bot.get("id")), "trial_started_channel": tenant_trial_fields.get("trial_started_channel")})
    return jsonify({"connected": True, "channel": "telegram", "bot": {"id": bot.get("id"), "username": bot.get("username"), "name": bot.get("first_name")}, "webhook_url": webhook_url, "pending_update_count": info.get("pending_update_count", 0)})


@platform_bp.get("/client/channels/<channel>/authorize")
@_require_tenant_roles("owner")
def authorize_client_channel(channel: str):
    try:
        channel = ensure_channel(channel)
        if channel not in {"instagram", "facebook", "tiktok", "x", "linkedin"}:
            return jsonify({"error": "Este canal não usa autorização OAuth."}), 400
        result = start_oauth(_db(), channel, _tenant_for_identity(_identity()), str((_identity() or {}).get("id") or ""))
        _audit("channel_oauth_started", _identity(), _tenant_for_identity(_identity()), {"channel": channel, "provider": result.get("provider")})
        return jsonify(result)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 404
    except RuntimeError as exc:
        return jsonify({"error": str(exc)}), 409


@platform_bp.get("/client/channels/<channel>/callback")
def authorize_client_channel_callback(channel: str):
    try:
        channel = ensure_channel(channel)
        result = complete_oauth(_db(), channel, str(request.args.get("code") or ""), str(request.args.get("state") or ""))
        tenant_id = str(result.get("tenant_id") or "")
        tenant = _tenant_data(tenant_id)
        central_account_id = central_account_id_for_tenant(tenant)
        registry = registry_status(_db(), central_account_id)
        if central_account_id and not is_paid_plan(tenant) and not registry.get("trial_consumed"):
            claimed, registry = claim_trial_for_account(_db(), central_account_id, tenant_id, channel, email=tenant.get("account_email") or tenant.get("email"))
            if claimed:
                tenant_ref = _db().collection("tenants").document(tenant_id)
                tenant_ref.set({**trial_fields_from_registry(registry, channel, result.get("external_account_id") or "oauth"), "central_account_id": central_account_id, "updated_at": _now()}, merge=True)
        _audit("channel_oauth_connected", {"role": "system"}, tenant_id, {"channel": channel, "external_account_id": result.get("external_account_id")})
        frontend = str(os.getenv("PUBLIC_APP_BASE_URL") or "https://app-negobotmoz.duckdns.org/plataforma").rstrip("/")
        return redirect(f"{frontend}/canais?oauth=success&channel={quote(channel)}")
    except (ValueError, RuntimeError) as exc:
        frontend = str(os.getenv("PUBLIC_APP_BASE_URL") or "https://app-negobotmoz.duckdns.org/plataforma").rstrip("/")
        return redirect(f"{frontend}/canais?oauth=error&channel={quote(str(channel))}")


@platform_bp.post("/client/channels/<channel>/disconnect")
@_require_tenant_roles("owner")
def disconnect_client_oauth_channel(channel: str):
    try:
        channel = ensure_channel(channel)
        if channel not in {"instagram", "facebook", "tiktok", "x", "linkedin"}:
            return jsonify({"error": "Este canal não usa desligamento OAuth."}), 400
        tenant_id = _tenant_for_identity(_identity())
        disconnect_oauth(_db(), channel, tenant_id)
        _audit("channel_oauth_disconnected", _identity(), tenant_id, {"channel": channel})
        return jsonify({"disconnected": True, "channel": channel, "status": "disabled"})
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 404


@platform_bp.post("/client/channels/telegram/disconnect")
@_require_tenant_roles("owner", "operator")
def disconnect_telegram():
    tenant_id = _tenant_for_identity(_identity())
    tenant_ref = _db().collection("tenants").document(tenant_id)
    tenant = tenant_ref.get().to_dict() or {}
    channels = dict(tenant.get("channels") or {}) if isinstance(tenant.get("channels"), dict) else {}
    telegram = dict(channels.get("telegram") or {})
    ciphertext = str(telegram.get("token_ciphertext") or "")
    if ciphertext:
        try:
            delete_webhook(decrypt_secret(ciphertext))
        except (TelegramApiError, SecretStoreError):
            pass
    channels["telegram"] = {"status": "disabled", "provider": "Telegram Bot API", "updated_at": _now()}
    tenant_ref.set({"channels": channels, "updated_at": _now()}, merge=True)
    _audit("telegram_webhook_disconnected", _identity(), tenant_id, {})
    return jsonify({"disconnected": True, "channel": "telegram"})


@platform_bp.patch("/client/channels/<channel>")
@_require_tenant_roles("owner")
def update_client_channel(channel: str):
    try:
        channel = ensure_channel(channel)
    except ValueError:
        return jsonify({"error": "Canal não suportado."}), 404
    payload = request.get_json(silent=True) or {}
    requested_status = str(payload.get("status") or "").strip().lower()
    if requested_status not in {"disabled", "not_configured"}:
        return jsonify({"error": "A ligação do canal será feita pelo fluxo seguro do fornecedor."}), 400
    tenant_id = _tenant_for_identity(_identity())
    reference = _db().collection("tenants").document(tenant_id)
    document = reference.get()
    if not document.exists:
        return jsonify({"error": "Tenant não encontrado."}), 404
    tenant = document.to_dict() or {}
    channels = dict(tenant.get("channels") or {}) if isinstance(tenant.get("channels"), dict) else {}
    config = dict(channels.get(channel) or {})
    config["status"] = requested_status
    config["updated_at"] = _now()
    channels[channel] = config
    reference.set({"channels": channels, "updated_at": _now()}, merge=True)
    _audit("client_channel_updated", _identity(), tenant_id, {"channel": channel, "status": requested_status})
    return jsonify({"updated": True, "channel": channel, "status": requested_status})
